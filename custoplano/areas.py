"""
Calculadora de R$/m² (uso interno da equipe Custo Plano).

O quadro de áreas é lido da prancha do projeto (PDF ou imagem) pela API da Anthropic, no servidor.
A chave fica só no servidor (variável ANTHROPIC_API_KEY no Render); sem ela a leitura automática
fica desligada e o quadro é preenchido à mão. Nada é usado antes de a pessoa conferir a tabela:
cada linha precisa fechar (computável + não computável = total) e a soma precisa bater com o total impresso.
"""
from __future__ import annotations

import base64
import datetime as dt
import io
import json
import os

MODELO = "claude-opus-5-5"
TIPOS = {"application/pdf": "document", "image/png": "image", "image/jpeg": "image", "image/webp": "image", "image/gif": "image"}
LIMITE_BYTES = 20 * 1024 * 1024
TOL_LINHA = 0.02   # m²: arredondamento de duas casas
TOL_TOTAL = 0.05


class ErroLeitura(Exception):
    """Mensagem para a pessoa; o código HTTP vai junto."""

    def __init__(self, msg: str, codigo: int = 400):
        super().__init__(msg)
        self.codigo = codigo


def ligada() -> bool:
    return bool(os.environ.get("ANTHROPIC_API_KEY"))


def _num():
    return {"anyOf": [{"type": "number"}, {"type": "null"}]}


ESQUEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["encontrado", "pavimentos", "total_computavel", "total_nao_computavel", "total_geral",
                 "privativa_total", "unidades_total", "terreno", "ca", "to", "observacoes"],
    "properties": {
        "encontrado": {"type": "boolean"},
        "pavimentos": {"type": "array", "items": {
            "type": "object",
            "additionalProperties": False,
            "required": ["nome", "computavel", "nao_computavel", "total", "privativa", "unidades"],
            "properties": {
                "nome": {"type": "string"},
                "computavel": _num(), "nao_computavel": _num(), "total": _num(),
                "privativa": _num(), "unidades": {"anyOf": [{"type": "integer"}, {"type": "null"}]},
            },
        }},
        "total_computavel": _num(), "total_nao_computavel": _num(), "total_geral": _num(),
        "privativa_total": _num(), "unidades_total": {"anyOf": [{"type": "integer"}, {"type": "null"}]},
        "terreno": _num(), "ca": _num(), "to": _num(),
        "observacoes": {"type": "string"},
    },
}

INSTRUCOES = """Esta é a prancha de um projeto de arquitetura. Encontre o QUADRO DE ÁREAS e transcreva os números exatamente como estão impressos.

Regras:
- Não calcule, não estime e não complete nada. Valor que não está impresso é null. Traço ("—" ou "-") também é null.
- Uma linha por pavimento, na ordem do quadro. Linha que agrupa pavimentos (ex.: "Tipo (1º ao 10º)") fica como uma linha só, com o nome impresso e os valores impressos nela.
- computavel, nao_computavel e total em m², com ponto decimal (2.381,70 vira 2381.70).
- total_computavel, total_nao_computavel e total_geral vêm da linha TOTAL impressa no quadro.
- privativa (área privativa) e unidades só se constarem no projeto; senão null.
- terreno (área do terreno em m²), ca (coeficiente de aproveitamento) e to (taxa de ocupação, como número: 0,38 vira 0.38; 38% vira 0.38) se estiverem impressos.
- encontrado = false se não houver quadro de áreas legível.
- Em observacoes, em português, diga em uma ou duas frases onde está o quadro e qualquer número de leitura duvidosa."""


def ler_quadro(conteudo: bytes, mime: str) -> dict:
    """Manda a prancha para o Claude e devolve o quadro de áreas como dicionário (ainda sem conferência)."""
    if not ligada():
        raise ErroLeitura("A leitura automática está desligada: falta a chave da Anthropic no servidor. Preencha o quadro à mão.", 503)
    tipo = TIPOS.get(mime)
    if not tipo:
        raise ErroLeitura("Envie a prancha em PDF ou em imagem (PNG, JPG ou WEBP).")
    if len(conteudo) > LIMITE_BYTES:
        raise ErroLeitura("Arquivo acima de 20 MB. Envie só a prancha com o quadro de áreas.")
    import anthropic

    bloco = {"type": tipo, "source": {"type": "base64", "media_type": mime, "data": base64.standard_b64encode(conteudo).decode()}}
    cliente = anthropic.Anthropic(timeout=180, max_retries=2)
    try:
        r = cliente.beta.messages.create(
            model=MODELO,
            max_tokens=16000,
            thinking={"type": "adaptive"},
            output_config={"format": {"type": "json_schema", "schema": ESQUEMA}},
            messages=[{"role": "user", "content": [bloco, {"type": "text", "text": INSTRUCOES}]}],
            betas=["server-side-fallback-2026-07-01"],
            fallbacks="default",
        )
    except anthropic.AuthenticationError:
        raise ErroLeitura("A chave da Anthropic no servidor foi recusada. Preencha o quadro à mão.", 503)
    except anthropic.RateLimitError:
        raise ErroLeitura("A leitura automática está ocupada agora. Tente de novo em um minuto.", 503)
    except anthropic.BadRequestError as e:
        print(f"Leitura do quadro recusada pela API: {e}", flush=True)
        raise ErroLeitura("Não consegui ler este arquivo. Tente uma imagem só do quadro de áreas.")
    except (anthropic.APIStatusError, anthropic.APIConnectionError) as e:
        print(f"Leitura do quadro falhou: {e}", flush=True)
        raise ErroLeitura("A leitura automática não respondeu. Tente de novo ou preencha à mão.", 502)
    u = r.usage
    print(f"Quadro de áreas lido: entrada {u.input_tokens} tokens, saída {u.output_tokens} tokens", flush=True)
    if r.stop_reason == "refusal":
        raise ErroLeitura("A leitura automática recusou este arquivo. Preencha o quadro à mão.")
    if r.stop_reason == "max_tokens":
        raise ErroLeitura("A leitura não terminou. Tente uma imagem só do quadro de áreas.", 502)
    texto = next((b.text for b in r.content if b.type == "text"), "")
    try:
        q = json.loads(texto)
    except ValueError:
        raise ErroLeitura("A leitura veio incompleta. Tente de novo ou preencha à mão.", 502)
    if not q.get("encontrado") or not q.get("pavimentos"):
        raise ErroLeitura("Não achei um quadro de áreas legível nesta prancha. " + (q.get("observacoes") or ""))
    return q


# ----------------------------------------------------------------- conferência
def _n(v) -> float | None:
    if v is None or v == "":
        return None
    return float(v)


def _soma(vals) -> float:
    return round(sum(v or 0 for v in vals), 2)


def validar(q: dict) -> dict:
    """Confere o quadro: cada linha fecha e as somas batem com os totais impressos."""
    linhas = []
    for p in q.get("pavimentos") or []:
        c, nc, t = _n(p.get("computavel")), _n(p.get("nao_computavel")), _n(p.get("total"))
        if t is None or (c is None and nc is None):
            linhas.append({"ok": False, "motivo": "faltam valores"})
        else:
            dif = round((c or 0) + (nc or 0) - t, 2)
            linhas.append({"ok": abs(dif) <= TOL_LINHA, "motivo": "" if abs(dif) <= TOL_LINHA else f"não fecha por {dif:+.2f} m²"})
    pv = q.get("pavimentos") or []
    somas = {"computavel": _soma(_n(p.get("computavel")) for p in pv),
             "nao_computavel": _soma(_n(p.get("nao_computavel")) for p in pv),
             "total": _soma(_n(p.get("total")) for p in pv)}
    totais = {}
    for k, impresso in (("computavel", "total_computavel"), ("nao_computavel", "total_nao_computavel"), ("total", "total_geral")):
        v = _n(q.get(impresso))
        totais[k] = {"impresso": v, "soma": somas[k], "ok": v is None or abs(somas[k] - v) <= TOL_TOTAL}
    ca = None
    terreno, ca_imp = _n(q.get("terreno")), _n(q.get("ca"))
    if terreno and ca_imp:
        calc = somas["computavel"] / terreno
        ca = {"impresso": ca_imp, "calculado": round(calc, 2), "ok": abs(calc - ca_imp) <= 0.01}
    return {"linhas": linhas, "totais": totais, "ca": ca,
            "ok": bool(pv) and all(l["ok"] for l in linhas) and all(t["ok"] for t in totais.values())}


def privativa(q: dict) -> float | None:
    v = _n(q.get("privativa_total"))
    if v:
        return v
    s = _soma(_n(p.get("privativa")) for p in q.get("pavimentos") or [])
    return s or None


def calcular(q: dict, venda: float | None, custo: float | None) -> dict:
    construida = _n(q.get("total_geral")) or validar(q)["totais"]["total"]["soma"]
    priv = privativa(q)
    r = {"construida": construida, "privativa": priv}
    for nome, valor in (("venda", venda), ("custo", custo)):
        if valor:
            r[nome] = {"construido": valor / construida if construida else None, "privativo": valor / priv if priv else None}
    return r


# ----------------------------------------------------------------- Excel
AZUL_ESCURO, CINZA, AZUL, VINHO, VERMELHO = "002455", "ECF0F1", "261FB3", "8E1616", "C62300"


def excel(q: dict, venda: float | None, custo: float | None, projeto: str, conferido_por: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    v = validar(q)
    res = calcular(q, venda, custo)
    wb = Workbook()
    ws = wb.active
    ws.title = "Quadro de áreas"
    fino = Side(style="thin", color="C8CED3")
    m2, reais = '#,##0.00 "m²"', '"R$" #,##0'
    ws["A1"] = "Quadro de áreas e R$/m²"
    ws["A1"].font = Font(size=16, bold=True, color=AZUL_ESCURO)
    ws["A2"] = projeto or "Projeto"
    ws["A2"].font = Font(size=12, color=AZUL_ESCURO)
    ws["A3"] = f"Áreas lidas da prancha e conferidas por {conferido_por} em {dt.date.today():%d/%m/%Y}."
    ws["A3"].font = Font(size=10, italic=True, color="515C66")

    cab = ["Pavimento", "Computável", "Não computável", "Total", "Privativa", "Unidades"]
    l0 = 5
    for j, h in enumerate(cab, start=1):
        c = ws.cell(l0, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL_ESCURO)
        c.alignment = Alignment(horizontal="left" if j == 1 else "right", vertical="center")
    ws.row_dimensions[l0].height = 22
    i = l0
    for k, p in enumerate(q.get("pavimentos") or []):
        i += 1
        vals = [p.get("nome"), _n(p.get("computavel")), _n(p.get("nao_computavel")), _n(p.get("total")), _n(p.get("privativa")), p.get("unidades")]
        for j, val in enumerate(vals, start=1):
            c = ws.cell(i, j, val)
            c.border = Border(bottom=fino)
            if 2 <= j <= 5:
                c.number_format = m2
            if k % 2:
                c.fill = PatternFill("solid", fgColor=CINZA)
            if not v["linhas"][k]["ok"]:
                c.font = Font(bold=True, color=VERMELHO)
    i += 1
    tot = ["Total", _n(q.get("total_computavel")) or v["totais"]["computavel"]["soma"],
           _n(q.get("total_nao_computavel")) or v["totais"]["nao_computavel"]["soma"],
           res["construida"], res["privativa"], q.get("unidades_total")]
    for j, val in enumerate(tot, start=1):
        c = ws.cell(i, j, val)
        c.font = Font(bold=True, color=AZUL_ESCURO)
        c.fill = PatternFill("solid", fgColor=CINZA)
        c.border = Border(top=Side(style="medium", color=AZUL_ESCURO))
        if 2 <= j <= 5:
            c.number_format = m2

    i += 2
    dados = [("Terreno", _n(q.get("terreno")), m2), ("Coeficiente de aproveitamento (CA)", _n(q.get("ca")), "0.00"),
             ("Taxa de ocupação (TO)", _n(q.get("to")), "0.00")]
    if v["ca"]:
        dados.append(("CA conferido (computável ÷ terreno)", v["ca"]["calculado"], "0.00"))
    for rot, val, fmt in dados:
        ws.cell(i, 1, rot).font = Font(color="515C66")
        c = ws.cell(i, 2, val if val is not None else "não consta")
        c.font = Font(bold=True)
        if val is not None:
            c.number_format = fmt
        i += 1

    i += 1
    c = ws.cell(i, 1, "R$/m²")
    c.font = Font(size=13, bold=True, color=AZUL_ESCURO)
    i += 1
    for j, h in enumerate(["", "Valor total", "R$/m² construído", "R$/m² privativo"], start=1):
        c = ws.cell(i, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL_ESCURO)
        c.alignment = Alignment(horizontal="left" if j == 1 else "right")
    for nome, rot, valor in (("venda", "Valor de venda", venda), ("custo", "Custo de obra", custo)):
        if not valor:
            continue
        i += 1
        ws.cell(i, 1, rot)
        ws.cell(i, 2, round(valor)).number_format = reais
        cc = ws.cell(i, 3, round(res[nome]["construido"]) if res[nome]["construido"] else None)
        cc.number_format, cc.font = reais, Font(bold=True, color=AZUL)
        if res[nome]["privativo"]:
            cp = ws.cell(i, 4, round(res[nome]["privativo"]))
            cp.number_format, cp.font = reais, Font(bold=True, color=AZUL)
        else:
            ws.cell(i, 4, "área privativa não consta").font = Font(italic=True, color=VINHO)
    if not res["privativa"]:
        i += 2
        ws.cell(i, 1, "A área privativa não consta no projeto, por isso o R$/m² privativo não foi calculado (não estimamos).").font = Font(italic=True, color=VINHO)

    for col, w in zip("ABCDEF", (38, 18, 18, 24, 16, 11)):
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
