"""
Plano de ataque (a "escadinha"): programação dos serviços pavimento por pavimento.

Lê a planilha no formato usado pela obra:
- um bloco com a linha "Serviço" (nomes dos serviços), "Fornecedor", "Ciclo" e,
  abaixo, uma linha por pavimento com a data de cada serviço e "ok" na coluna
  seguinte quando o serviço foi executado;
- blocos seguintes marcados "Linha de Base ..." no mesmo formato (sem o "ok");
- opcionais: DATA BASE, CONTROLE DE PRAZO - DATAS MARCO, TÉRMINO DA OBRA e
  META TRIMESTRAL.
"""
from __future__ import annotations

import datetime as dt
import io
import re
import unicodedata

import openpyxl


class ErroPlano(Exception):
    pass


def norm(v) -> str:
    """Texto em minúsculas, sem acentos e sem espaços extras."""
    if v is None:
        return ""
    s = unicodedata.normalize("NFKD", str(v)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower().rstrip(":")


def data(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # número de série do Excel
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
    if isinstance(v, str):
        m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", v)
        if m:
            d, mes, a = int(m[1]), int(m[2]), int(m[3])
            try:
                return dt.date(a + 2000 if a < 100 else a, mes, d)
            except ValueError:
                return None
    return None


def iso(d: dt.date | None) -> str | None:
    return d.isoformat() if d else None


def numero(v) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        m = re.search(r"-?\d+(?:[.,]\d+)?", v)
        if m:
            return float(m[0].replace(",", "."))
    return None


def num_pavimento(rotulo, i: int) -> float:
    t = norm(rotulo)
    if "terreo" in t:
        return 0
    if "cobertura" in t or "atico" in t:
        return 1000
    n = numero(rotulo)
    return n if n is not None else -1000 - i


class Planilha:
    def __init__(self, linhas: list[tuple]):
        self.L = [list(r) for r in linhas]
        self.largura = max((len(r) for r in self.L), default=0)
        for r in self.L:
            r.extend([None] * (self.largura - len(r)))

    def v(self, r: int, c: int):
        if 0 <= r < len(self.L) and 0 <= c < self.largura:
            return self.L[r][c]
        return None

    def achar(self, *textos, linhas=None):
        """Posições (linha, coluna) das células cujo texto normalizado é um dos textos."""
        alvo = {norm(t) for t in textos}
        faixa = linhas if linhas is not None else range(len(self.L))
        return [(r, c) for r in faixa for c in range(self.largura) if norm(self.L[r][c]) in alvo]


def escolher_aba(wb):
    for ws in wb.worksheets:
        if "escad" in norm(ws.title) or "ataque" in norm(ws.title):
            return ws
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=60, max_col=3, values_only=True):
            if any(norm(v) == "servico" for v in row):
                return ws
    raise ErroPlano("Não encontrei a linha “Serviço” com os serviços em nenhuma aba da planilha.")


def ler_bloco(P: Planilha, r0: int, fim: int, com_status: bool) -> dict:
    """Bloco que começa na linha 'Serviço' (r0) e vai até antes de 'fim'."""
    servicos = []
    for c in range(1, P.largura):
        nome = P.v(r0, c)
        if isinstance(nome, str) and nome.strip():
            servicos.append({"nome": nome.strip(), "col": c})
    if not servicos:
        return {}
    cols = {s["col"] for s in servicos}

    fornec, ciclo, pavs = {}, {}, []
    vazias = 0
    for r in range(r0 + 1, fim):
        rot = P.v(r, 0)
        t = norm(rot)
        if t == "fornecedor":
            fornec = {s["col"]: P.v(r, s["col"]) for s in servicos}
            continue
        if t == "ciclo":
            ciclo = {s["col"]: P.v(r, s["col"]) for s in servicos}
            continue
        if rot is None or t == "":
            vazias += 1
            if pavs and vazias >= 2:
                break
            continue
        if "linha de base" in t or t == "servico" or t == "legenda":
            break
        vazias = 0
        linha_datas, linha_ok, alguma = [], [], False
        for s in servicos:
            d = data(P.v(r, s["col"]))
            alguma = alguma or d is not None
            linha_datas.append(iso(d))
            ok = False
            if com_status and s["col"] + 1 not in cols:
                ok = norm(P.v(r, s["col"] + 1)) in ("ok", "x", "sim", "executado", "concluido")
            linha_ok.append(ok)
        if not alguma:
            continue
        rotulo = str(rot).strip()
        if isinstance(rot, (int, float)):
            rotulo = f"Pav. {int(rot)}"
        pavs.append({"rotulo": rotulo, "n": num_pavimento(rot, len(pavs)), "datas": linha_datas, "ok": linha_ok})

    if not pavs:
        return {}
    # só serviços que têm alguma data
    usados = [i for i in range(len(servicos)) if any(p["datas"][i] for p in pavs)]

    def ciclo_txt(v):
        if v is None or (isinstance(v, str) and v.startswith("#")):
            return None
        n = numero(v)
        return n if n is not None else None

    return {
        "servicos": [{
            "nome": servicos[i]["nome"],
            "fornecedor": str(fornec.get(servicos[i]["col"]) or "").strip(),
            "ciclo": ciclo_txt(ciclo.get(servicos[i]["col"])),
        } for i in usados],
        "pavimentos": [{
            "rotulo": p["rotulo"], "n": p["n"],
            "datas": [p["datas"][i] for i in usados],
            "ok": [p["ok"][i] for i in usados] if com_status else None,
        } for p in pavs],
    }


def ler_marcos(P: Planilha) -> list[dict]:
    for r, c in P.achar("Atividade"):
        cab = {norm(P.v(r, k)): k for k in range(P.largura) if P.v(r, k) is not None}
        if "linha de base" not in cab:
            continue
        col_lb, col_ant, col_at = cab.get("linha de base"), cab.get("mes anterior"), cab.get("atual")
        out = []
        for rr in range(r + 1, min(r + 15, len(P.L))):
            nome = P.v(rr, c)
            if not isinstance(nome, str) or not nome.strip():
                break
            lb, ant, at = (data(P.v(rr, k)) if k is not None else None for k in (col_lb, col_ant, col_at))
            out.append({
                "nome": nome.strip(), "lb": iso(lb), "anterior": iso(ant), "atual": iso(at),
                "var_lb": (at - lb).days if at and lb else None,
                "var_mes": (at - ant).days if at and ant else None,
            })
        return out
    return []


def ler_termino(P: Planilha) -> dict | None:
    for r, c in P.achar("Término da obra"):
        for rr in range(r + 1, min(r + 4, len(P.L))):
            cab = {norm(P.v(rr, k)): k for k in range(c - 3, P.largura) if P.v(rr, k) is not None}
            if "lb" in cab:
                lb = data(P.v(rr + 1, cab["lb"]))
                real = None
                for chave in ("real", "atual", "tendencia"):
                    if chave in cab:
                        real = data(P.v(rr + 1, cab[chave]))
                        break
                if lb:
                    return {"lb": iso(lb), "atual": iso(real), "var": (real - lb).days if real else None}
    return None


def ler_metas(P: Planilha) -> list[dict]:
    for r, c in P.achar("Mês"):
        cab = {norm(P.v(r, k)): k for k in range(c, min(c + 12, P.largura)) if P.v(r, k) is not None}
        if "real" not in cab or "meta" not in cab:
            continue
        out = []
        for rr in range(r + 1, min(r + 20, len(P.L))):
            mes = data(P.v(rr, c))
            if not mes:
                break
            real, meta = numero(P.v(rr, cab["real"])), numero(P.v(rr, cab["meta"]))
            out.append({"mes": iso(mes), "real": real, "meta": meta,
                        "idp": (real / meta) if real is not None and meta else None})
        return out
    return []


def ler_plano(nome_arquivo: str, conteudo: bytes) -> dict:
    if not re.search(r"\.xls[xm]$", nome_arquivo or "", re.I):
        raise ErroPlano("Envie a planilha em Excel (.xlsx ou .xlsm).")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    except Exception:
        raise ErroPlano("Não consegui abrir a planilha. Confira se é um arquivo .xlsx válido.")
    try:
        ws = escolher_aba(wb)
        P = Planilha(list(ws.iter_rows(min_row=1, max_row=600, max_col=200, values_only=True)))
        aba = ws.title
    finally:
        wb.close()

    cabecalhos = [r for r, c in P.achar("Serviço") if c == 0] or [r for r, _ in P.achar("Serviço")]
    if not cabecalhos:
        raise ErroPlano("Não encontrei a linha “Serviço” com os nomes dos serviços.")

    blocos = []
    for i, r0 in enumerate(cabecalhos):
        fim = cabecalhos[i + 1] if i + 1 < len(cabecalhos) else len(P.L)
        rotulo = None
        for rr in range(max(0, r0 - 5), r0):
            for c in range(min(8, P.largura)):
                if "linha de base" in norm(P.v(rr, c)):
                    rotulo = str(P.v(rr, c)).strip()
        blocos.append((r0, fim, rotulo))

    atual_def = next((b for b in blocos if not b[2]), blocos[0])
    atual = ler_bloco(P, atual_def[0], atual_def[1], com_status=True)
    if not atual:
        raise ErroPlano("Encontrei a linha “Serviço”, mas nenhum pavimento com datas abaixo dela.")

    nomes_atual = {norm(s["nome"]) for s in atual["servicos"]}
    base, melhor = None, -1
    for r0, fim, rotulo in blocos:
        if (r0, fim, rotulo) == atual_def or not rotulo:
            continue
        b = ler_bloco(P, r0, fim, com_status=False)
        if not b:
            continue
        comum = len(nomes_atual & {norm(s["nome"]) for s in b["servicos"]})
        if comum >= melhor:  # empate: fica com o último bloco (mais recente)
            base, melhor = dict(b, rotulo=rotulo), comum

    data_base = None
    for r, c in P.achar("Data base"):
        for rr, cc in ((r + 1, c), (r, c + 1), (r, c + 2), (r + 1, c + 1)):
            data_base = data(P.v(rr, cc))
            if data_base:
                break
        if data_base:
            break

    titulo = None
    for r, c in P.achar("Escadinha", "Plano de ataque", linhas=range(min(6, len(P.L)))):
        t = P.v(r + 1, c)
        if isinstance(t, str) and t.strip():
            titulo = t.strip()
            break

    return {
        "arquivo": nome_arquivo,
        "aba": aba,
        "titulo": titulo or re.sub(r"\.[^.]+$", "", nome_arquivo),
        "data_base": iso(data_base),
        "atual": atual,
        "base": base,
        "marcos": ler_marcos(P),
        "termino": ler_termino(P),
        "metas": ler_metas(P),
    }


# ------------------------------------------------- plano a partir do cronograma
_PAV = [
    (re.compile(r"\b(?:pav(?:to|imento)?\.?\s*)?tipo\s*(\d{1,3})\b"), lambda m: float(m[1])),
    (re.compile(r"\b(\d{1,3})\s*o?\s*(?:pav(?:to|imento)?|andar|pvto)\b\.?"), lambda m: float(m[1])),
    (re.compile(r"\b(?:pav(?:to|imento)?|andar|pvto|piso)\.?\s*(\d{1,3})\b"), lambda m: float(m[1])),
    (re.compile(r"\b(\d{1,3})\s*o?\s*subsolo\b|\bsubsolo\s*(\d{0,2})\b"), lambda m: -float(m[1] or m[2] or 1)),
    (re.compile(r"\bterreo\b"), lambda m: 0.0),
    (re.compile(r"\bmezanino\b"), lambda m: 0.5),
    (re.compile(r"\bcobertura\b"), lambda m: 1000.0),
    (re.compile(r"\batico\b"), lambda m: 1001.0),
    (re.compile(r"\bbarrilete\b"), lambda m: 1002.0),
    (re.compile(r"\bcasa de maquinas\b"), lambda m: 1003.0),
]


def _norm_igual(nome: str) -> str:
    """Minúsculas e sem acentos, com o mesmo número de caracteres do original."""
    out = []
    for ch in nome:
        d = unicodedata.normalize("NFKD", ch).encode("ascii", "ignore").decode()
        out.append(d[0].lower() if d else " ")
    return "".join(out)


def achar_pavimento(nome: str):
    """Procura um pavimento no nome da tarefa. Devolve (número, trecho encontrado) ou None."""
    t = _norm_igual(nome)
    for rx, f in _PAV:
        m = rx.search(t)
        if m:
            return f(m), m
    return None


def _rotulo_pav(n: float) -> str:
    if n == 0:
        return "Térreo"
    if n == 0.5:
        return "Mezanino"
    if n < 0:
        return f"Subsolo {int(-n)}"
    return {1000: "Cobertura", 1001: "Ático", 1002: "Barrilete", 1003: "Casa de máquinas"}.get(n, f"{int(n)}º pav.")


def _sem_pavimento(nome: str, trecho) -> str:
    """Nome do serviço sem o pavimento (ex.: 'Alvenaria – 3º pav.' -> 'Alvenaria')."""
    a, b = trecho.span()
    s = (nome[:a] + " " + nome[b:]).strip()
    s = re.sub(r"\(\s*\)|\[\s*\]", " ", s)
    s = re.sub(r"^[\s\-–—:|/.,]+|[\s\-–—:|/.,(]+$", "", s)
    s = re.sub(r"\s{2,}", " ", s)
    return s.strip()


def plano_do_projeto(p, atual: dict | None = None, data_status: dt.datetime | None = None) -> dict | None:
    """Monta o plano de ataque (serviço × pavimento) a partir das tarefas do cronograma.

    Reconhece as duas formas usuais de montar a EAP: Serviço > Pavimento ou Pavimento > Serviço
    (o pavimento pode estar no nome da tarefa, ex.: 'Alvenaria 3º pav.', ou numa tarefa-resumo acima).
    """
    atual = atual or {}
    celulas: dict[tuple[str, float], dict] = {}
    nomes: dict[str, str] = {}
    ordem_serv: dict[str, dt.datetime] = {}
    for t in p.tarefas:
        if t.filhos or t.marco or not t.inicio:
            continue
        achado = achar_pavimento(t.nome)
        servico = None
        if achado:
            n, trecho = achado
            servico = _sem_pavimento(t.nome, trecho) or (t.pai.nome if t.pai else "")
        else:
            a = t.pai
            while a is not None and not (achado := achar_pavimento(a.nome)):
                a = a.pai
            if not achado:
                continue
            n = achado[0]
            servico = t.nome
        if not servico:
            continue
        chave = norm(servico)
        nomes.setdefault(chave, servico)
        c = celulas.setdefault((chave, n), {"ini": t.inicio, "lb": t.lb_inicio, "ok": True})
        c["ini"] = min(c["ini"], t.inicio)
        if t.lb_inicio:
            c["lb"] = min(c["lb"], t.lb_inicio) if c["lb"] else t.lb_inicio
        c["ok"] = c["ok"] and float(atual.get(t.uid, t.pct) or 0) >= 99.5
        ordem_serv[chave] = min(ordem_serv.get(chave, t.inicio), t.inicio)

    pavs = sorted({n for _, n in celulas})
    # serviço que aparece em um pavimento só não forma escadinha
    servs = [s for s in sorted(ordem_serv, key=lambda s: ordem_serv[s])
             if sum(1 for (k, _) in celulas if k == s) >= 2]
    if len(pavs) < 2 or not servs:
        return None

    def bloco(campo, com_status):
        return {
            "servicos": [{"nome": nomes[s], "fornecedor": "", "ciclo": None} for s in servs],
            "pavimentos": [{
                "rotulo": _rotulo_pav(n), "n": n,
                "datas": [iso(celulas[(s, n)][campo].date()) if (s, n) in celulas and celulas[(s, n)][campo] else None for s in servs],
                "ok": [bool(celulas.get((s, n), {}).get("ok")) for s in servs] if com_status else None,
            } for n in pavs],
        }

    atual_b = bloco("ini", True)
    tem_lb = any(c["lb"] for c in celulas.values())
    marcos = []
    for t in p.tarefas:
        if t.marco and t.termino:
            lb = t.lb_termino.date() if t.lb_termino else None
            at = t.termino.date()
            marcos.append({"nome": t.nome, "lb": iso(lb), "anterior": None, "atual": iso(at),
                           "var_lb": (at - lb).days if lb else None, "var_mes": None})
    raiz = [t for t in p.tarefas if t.pai is None]
    fim = max((t.termino for t in raiz if t.termino), default=None)
    fim_lb = max((t.lb_termino for t in raiz if t.lb_termino), default=None)
    termino = ({"lb": iso(fim_lb.date()), "atual": iso(fim.date()), "var": (fim.date() - fim_lb.date()).days}
               if fim and fim_lb else None)
    return {
        "arquivo": p.arquivo, "aba": None, "origem": "project",
        "titulo": p.titulo,
        "data_base": iso((data_status or p.data_status or dt.datetime.now()).date()),
        "atual": atual_b,
        "base": dict(bloco("lb", False), rotulo="Linha de base do Project") if tem_lb else None,
        "marcos": marcos[:12],
        "termino": termino,
        "metas": [],
    }
