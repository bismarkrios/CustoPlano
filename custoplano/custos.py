"""
Custos da obra: orçamento com composições, cronograma financeiro e orçado × realizado.

Planilha (.xlsx) com as abas:
- Orçamento:   Código | Descrição | Unidade | Quantidade | Composição | Custo unitário | Tarefa (EAP) | Início | Término
               linhas sem quantidade e sem custo são títulos de etapa (ex.: "1.2  Fundações")
- Composições: Composição | Descrição da composição | Unidade | Tipo | Insumo | Unidade do insumo | Coeficiente | Preço unitário
               (uma linha por insumo; o código da composição pode ficar em branco nas linhas seguintes)
- Realizado:   Data | Código (item ou etapa) | Descrição | Fornecedor | Tipo | Valor          (opcional)
- Parâmetros:  BDI (%)                                                                         (opcional)

O controle de custo usa o custo direto (o BDI entra só no preço de venda).
O cálculo liga cada item a uma tarefa do cronograma (coluna Tarefa, ou pelo código/nome da etapa)
para distribuir o custo no tempo pela linha de base e medir o valor agregado pelo % físico.
"""
from __future__ import annotations

import datetime as dt
import io
import re
import unicodedata

import openpyxl


class ErroCustos(Exception):
    pass


TIPOS = {"mat": "Material", "mo": "Mão de obra", "eq": "Equipamento", "sv": "Serviço", "out": "Outros"}


def norm(v) -> str:
    if v is None:
        return ""
    s = unicodedata.normalize("NFKD", str(v)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower().rstrip(":").strip()


def tipo_de(v) -> str:
    t = norm(v)
    if not t:
        return "out"
    if t.startswith("mat") or t in ("m", "insumo"):
        return "mat"
    if "obra" in t or t in ("mo", "m.o.", "m.o", "mao-de-obra") or t.startswith("mao"):
        return "mo"
    if t.startswith("equip") or t in ("eq", "e"):
        return "eq"
    if t.startswith("serv") or t.startswith("empreit") or t in ("sv", "s", "terceiro", "terceiros"):
        return "sv"
    return "out"


def numero(v) -> float | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("R$", "").replace("%", "").strip()
    if not s:
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def data(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
    if isinstance(v, str):
        m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", v)
        if m:
            a = int(m[3])
            try:
                return dt.date(a + 2000 if a < 100 else a, int(m[2]), int(m[1]))
            except ValueError:
                return None
        m = re.search(r"(\d{4})-(\d{2})-(\d{2})", v)
        if m:
            return dt.date(int(m[1]), int(m[2]), int(m[3]))
    return None


def texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ------------------------------------------------------------------ leitura
def _cabecalho(linhas, sinonimos: dict, obrigatorios: tuple, max_linhas=30):
    """Acha a linha de cabeçalho e devolve (índice da linha, {campo: coluna})."""
    for r, linha in enumerate(linhas[:max_linhas]):
        nomes = [norm(v) for v in linha]
        mapa = {}
        for campo, opcoes in sinonimos.items():
            for c, n in enumerate(nomes):
                if n and c not in mapa.values() and any(n == o or (len(o) > 3 and n.startswith(o)) for o in opcoes):
                    mapa[campo] = c
                    break
        if all(k in mapa for k in obrigatorios):
            return r, mapa
    return None, None


def _aba(wb, *chaves):
    for ws in wb.worksheets:
        t = norm(ws.title)
        if any(k in t for k in chaves):
            return ws
    return None


def _linhas(ws, max_row=5000, max_col=40):
    return [list(r) for r in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True)]


SIN_ORC = {
    "codigo": ("codigo", "cod", "item", "cod."),
    "descricao": ("descricao", "descricao do servico", "servico", "discriminacao"),
    "unidade": ("unidade", "un", "und", "unid"),
    "qtd": ("quantidade", "qtd", "quant", "qtde"),
    "comp": ("composicao", "cod composicao", "codigo da composicao", "codigo composicao", "comp"),
    "preco": ("custo unitario", "preco unitario", "valor unitario", "unitario", "pu", "preco unit"),
    "total": ("custo total", "valor total", "total", "preco total"),
    "tarefa": ("tarefa (eap)", "tarefa", "eap da tarefa", "eap tarefa", "eap", "atividade"),
    "inicio": ("inicio", "data inicio", "data de inicio"),
    "termino": ("termino", "fim", "data termino", "data de termino"),
}
SIN_COMP = {
    "comp": ("composicao", "codigo da composicao", "cod composicao", "codigo"),
    "desc_comp": ("descricao da composicao", "descricao composicao"),
    "un_comp": ("unidade da composicao", "unidade composicao"),
    "tipo": ("tipo", "classe", "categoria"),
    "insumo": ("insumo", "descricao do insumo", "descricao insumo"),
    "un": ("unidade do insumo", "unidade insumo", "unidade", "un"),
    "coef": ("coeficiente", "coef", "consumo", "indice"),
    "preco": ("preco unitario", "custo unitario", "valor unitario", "preco", "pu"),
}
SIN_REAL = {
    "data": ("data", "data do pagamento", "data de pagamento", "competencia"),
    "codigo": ("codigo", "item", "etapa", "cod"),
    "descricao": ("descricao", "historico"),
    "fornecedor": ("fornecedor", "favorecido", "empresa"),
    "tipo": ("tipo", "classe", "categoria"),
    "valor": ("valor", "valor pago", "total"),
}


# ------------------------------- orçamento no padrão Custo Plano (árvore UA, gerado por planilha_custoplano.py)
def _cabecalho_ua(L) -> tuple[int, dict] | None:
    """Linha do cabeçalho UA | CÓDIGO | NÍVEL | DESCRIÇÃO | QTDE | UNID. | UNITÁRIO | VALOR e a coluna de cada um."""
    for r, linha in enumerate(L):
        m = {}
        for c, v in enumerate(linha):
            n = norm(v).rstrip(".")
            for chave, nomes in (("ua", ("ua",)), ("sigla", ("codigo",)), ("nivel", ("nivel",)), ("desc", ("descricao",)),
                                 ("qtd", ("qtde", "quantidade")), ("un", ("unid", "unidade")),
                                 ("unit", ("unitario (r$)", "unitario")), ("valor", ("valor (r$)", "valor"))):
                if n in nomes and chave not in m:
                    m[chave] = c
        if {"ua", "nivel", "desc", "qtd", "valor"} <= set(m):
            return r, m
    return None


def _arred(v: float, casas: int = 0) -> float:
    """ROUND do Excel (metade para longe do zero), não o arredondamento bancário do Python."""
    from decimal import ROUND_HALF_UP, Decimal
    return float(Decimal(repr(v)).quantize(Decimal(1).scaleb(-int(casas)), rounding=ROUND_HALF_UP))


class _Formulas:
    """Calcula as fórmulas simples da planilha UA (referências, SUM, ROUND e as quatro operações), para ler
    um arquivo gerado pelo Python que nunca foi aberto no Excel e por isso não tem os valores guardados."""
    REF = re.compile(r"(?:'?([A-Za-zÀ-ú ]+)'?!)?\$?([A-Z]{1,2})\$?(\d+)")

    def __init__(self, wb_f, wb_v, aba: str):
        self.f, self.v, self.aba, self.memo = wb_f, wb_v, aba, {}

    def valor(self, aba: str, col: str, lin: int, prof: int = 0):
        chave = (aba, col, lin)
        if chave in self.memo:
            return self.memo[chave]
        if prof > 200 or aba not in self.f.sheetnames:
            raise ValueError("referência")
        guardado = self.v[aba][f"{col}{lin}"].value
        bruto = self.f[aba][f"{col}{lin}"].value
        if isinstance(guardado, (int, float)) and not isinstance(guardado, bool):
            r = float(guardado)
        elif bruto is None or bruto == "":
            r = 0.0
        elif isinstance(bruto, (int, float)) and not isinstance(bruto, bool):
            r = float(bruto)
        elif isinstance(bruto, str) and bruto.startswith("="):
            r = self.calcular(aba, bruto[1:], prof)
        else:
            n = numero(bruto)
            if n is None:
                raise ValueError("texto")
            r = n
        self.memo[chave] = r
        return r

    def calcular(self, aba: str, expr: str, prof: int) -> float:
        from openpyxl.utils import column_index_from_string, get_column_letter

        def soma(m):
            a, b = m[1], m[2]
            ra, rb = self.REF.fullmatch(a), self.REF.fullmatch(b)
            if not ra or not rb:
                raise ValueError("intervalo")
            folha = ra[1] or aba
            c0, c1 = sorted((column_index_from_string(ra[2]), column_index_from_string(rb[2])))
            l0, l1 = sorted((int(ra[3]), int(rb[3])))
            return "(" + "+".join(repr(self.valor(folha, get_column_letter(c), lin, prof + 1))
                                  for c in range(c0, c1 + 1) for lin in range(l0, l1 + 1)) + ")"
        e = re.sub(r"(?i)\bSUM\(\s*([^:()]+?)\s*:\s*([^:()]+?)\s*\)", soma, expr)
        e = re.sub(r"(?i)\bROUND\(", "R(", e)
        e = self.REF.sub(lambda m: repr(self.valor((m[1] or aba).strip(), m[2], int(m[3]), prof + 1)), e)
        if not re.fullmatch(r"[0-9eE.+\-*/(), R]*", e):
            raise ValueError("fórmula")
        return float(eval(e, {"__builtins__": {}}, {"R": _arred}))  # noqa: S307 - só números e operadores


def _ler_padrao(nome_arquivo: str, conteudo: bytes) -> dict:
    """Orçamento analítico do padrão Custo Plano: UA-2 viram etapas, UA-3 (GSA/GSE) viram subetapas e os
    serviços UA-4 viram itens com o código UA, que liga cada item às tarefas do cronograma padrão."""
    wb_f = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=False)
    wb_v = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
    aba = next(w.title for w in wb_f.worksheets if norm(w.title) == "planilha de orcamento")
    L = _linhas(wb_f[aba])
    r0, m = _cabecalho_ua(L)
    from openpyxl.utils import get_column_letter
    calc = _Formulas(wb_f, wb_v, aba)
    col = {k: get_column_letter(c + 1) for k, c in m.items()}
    g = lambda linha, k: linha[m[k]] if k in m and m[k] < len(linha) else None  # noqa: E731
    itens, ua2, ua3, n, sem_valor = [], "", None, 0, []
    for i, linha in enumerate(L[r0 + 1:], start=r0 + 2):
        nivel = norm(g(linha, "nivel")).upper()
        desc = texto(g(linha, "desc"))
        if nivel in ("UA-0", "UA-1"):
            continue
        if nivel == "UA-2":
            ua2 = texto(g(linha, "ua")).split(" ")[0] or ua2
            ua3 = None
            itens.append({"codigo": ua2, "descricao": desc, "titulo": True, "unidade": "", "qtd": 0.0, "comp": "",
                          "preco": None, "total_planilha": None, "tarefa": "", "inicio": None, "termino": None})
        elif nivel == "UA-3":
            sig = texto(g(linha, "sigla")) or texto(g(linha, "ua"))
            ua3, n = sig, 0
            itens.append({"codigo": f"{ua2}.{sig}", "descricao": desc, "titulo": True, "unidade": "", "qtd": 0.0,
                          "comp": "", "preco": None, "total_planilha": None, "tarefa": "", "inicio": None, "termino": None})
        elif nivel == "UA-4" and desc:
            n += 1
            try:
                valor = calc.valor(aba, col["valor"], i)
                qtd = calc.valor(aba, col["qtd"], i)
            except (ValueError, KeyError, SyntaxError, ZeroDivisionError):
                sem_valor.append(desc)
                continue
            itens.append({"codigo": f"{ua2}.{ua3 or 'SEM'}.{n}", "descricao": desc, "titulo": False,
                          "unidade": texto(g(linha, "un")), "qtd": qtd if qtd else 1.0, "comp": "",
                          "preco": (valor / qtd) if qtd else valor, "total_planilha": round(valor, 2),
                          "tarefa": "", "ua": ua3, "inicio": None, "termino": None})
    wb_f.close()
    wb_v.close()
    if not any(not x["titulo"] for x in itens):
        raise ErroCustos("Na Planilha de Orçamento não encontrei nenhum serviço (linhas UA-4) com valor.")
    avisos = []
    if sem_valor:
        avisos.append(f"{len(sem_valor)} serviços ficaram de fora porque o valor não pôde ser calculado: "
                      + ", ".join(sem_valor[:5]) + ("…" if len(sem_valor) > 5 else "")
                      + ". Abra a planilha no Excel, salve e envie de novo.")
    # o BDI da planilha padrão é só de referência (o orçamento é de custo): não entra
    return {"arquivo": nome_arquivo, "origem": "padrao", "bdi": 0.0, "itens": itens, "composicoes": {},
            "realizado": [], "avisos": avisos}


def ler_custos(nome_arquivo: str, conteudo: bytes) -> dict:
    if not re.search(r"\.xls[xm]$", nome_arquivo or "", re.I):
        raise ErroCustos("Envie o orçamento em Excel (.xlsx). Baixe o modelo na tela Custos se precisar.")
    try:
        if not _zip_seguro(conteudo):
            raise ErroCustos("A planilha é grande demais para ser lida.")
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    except Exception:
        raise ErroCustos("Não consegui abrir a planilha. Confira se é um arquivo .xlsx válido.")
    try:
        ws_ua = next((w for w in wb.worksheets if norm(w.title) == "planilha de orcamento"), None)
        if ws_ua is not None and _cabecalho_ua(_linhas(ws_ua, 40, 20)) is not None:
            return _ler_padrao(nome_arquivo, conteudo)
        ws_orc = _aba(wb, "orcament") or wb.worksheets[0]
        L = _linhas(ws_orc)
        r0, m = _cabecalho(L, SIN_ORC, ("codigo", "descricao"))
        if r0 is None or ("qtd" not in m and "total" not in m):
            raise ErroCustos("Na aba Orçamento não encontrei o cabeçalho com Código, Descrição e Quantidade.")

        bdi = None
        ws_par = _aba(wb, "parametr")
        for linhas in ([_linhas(ws_par, 50, 10)] if ws_par else []) + [L[:r0]]:
            for linha in linhas:
                for c, v in enumerate(linha):
                    if "bdi" in norm(v):
                        for w in linha[c + 1:]:
                            n = numero(w)
                            if n is not None:
                                bdi = n / 100 if n > 1 else n
                                break
                    if bdi is not None:
                        break
                if bdi is not None:
                    break
            if bdi is not None:
                break

        g = lambda linha, k: linha[m[k]] if k in m and m[k] < len(linha) else None  # noqa: E731
        itens = []
        for linha in L[r0 + 1:]:
            cod, desc = texto(g(linha, "codigo")), texto(g(linha, "descricao"))
            if not cod and not desc:
                continue
            qtd, preco, total = numero(g(linha, "qtd")), numero(g(linha, "preco")), numero(g(linha, "total"))
            comp = texto(g(linha, "comp"))
            titulo = qtd is None and preco is None and not comp and not texto(g(linha, "unidade"))
            if norm(desc).startswith("total"):
                continue
            itens.append({
                "codigo": cod, "descricao": desc, "titulo": bool(titulo),
                "unidade": texto(g(linha, "unidade")), "qtd": qtd if qtd is not None else (1.0 if total and not titulo else 0.0),
                "comp": comp, "preco": preco if preco is not None else (total if (total and qtd in (None, 0)) else None),
                "total_planilha": total,
                "tarefa": texto(g(linha, "tarefa")),
                "inicio": (data(g(linha, "inicio")) or None) and data(g(linha, "inicio")).isoformat(),
                "termino": (data(g(linha, "termino")) or None) and data(g(linha, "termino")).isoformat(),
            })
        if not any(not i["titulo"] for i in itens):
            raise ErroCustos("A aba Orçamento não tem nenhum item com quantidade ou custo.")

        comps = {}
        ws_c = _aba(wb, "composic")
        if ws_c:
            LC = _linhas(ws_c)
            rc, mc = _cabecalho(LC, SIN_COMP, ("comp", "coef", "preco"))
            if rc is not None:
                gc = lambda linha, k: linha[mc[k]] if k in mc and mc[k] < len(linha) else None  # noqa: E731
                atual = None
                for linha in LC[rc + 1:]:
                    cod = texto(gc(linha, "comp"))
                    if cod:
                        atual = comps.setdefault(cod, {"codigo": cod, "descricao": texto(gc(linha, "desc_comp")),
                                                       "unidade": texto(gc(linha, "un_comp")), "insumos": []})
                        if not atual["descricao"]:
                            atual["descricao"] = texto(gc(linha, "desc_comp"))
                    coef, preco = numero(gc(linha, "coef")), numero(gc(linha, "preco"))
                    if atual is None or coef is None or preco is None:
                        continue
                    atual["insumos"].append({"tipo": tipo_de(gc(linha, "tipo")), "descricao": texto(gc(linha, "insumo")),
                                             "unidade": texto(gc(linha, "un")), "coef": coef, "preco": preco})

        realizado = []
        ws_r = _aba(wb, "realizad", "lancament", "despesa", "pagament")
        if ws_r:
            LR = _linhas(ws_r)
            rr, mr = _cabecalho(LR, SIN_REAL, ("data", "valor"))
            if rr is not None:
                gr = lambda linha, k: linha[mr[k]] if k in mr and mr[k] < len(linha) else None  # noqa: E731
                for linha in LR[rr + 1:]:
                    d, v = data(gr(linha, "data")), numero(gr(linha, "valor"))
                    if not d or v is None:
                        continue
                    realizado.append({"data": d.isoformat(), "codigo": texto(gr(linha, "codigo")),
                                      "descricao": texto(gr(linha, "descricao")), "fornecedor": texto(gr(linha, "fornecedor")),
                                      "tipo": tipo_de(gr(linha, "tipo")), "valor": v, "origem": "planilha"})
    finally:
        wb.close()

    faltando = sorted({i["comp"] for i in itens if i["comp"] and i["comp"] not in comps})
    return {"arquivo": nome_arquivo, "bdi": bdi or 0.0, "itens": itens, "composicoes": comps,
            "realizado": realizado, "avisos": ([f"Composições citadas no orçamento e não encontradas: {', '.join(faltando[:8])}"
                                                + ("…" if len(faltando) > 8 else "")] if faltando else [])}


# --------------------------------------------------------------- análise
def _mes(d: dt.date) -> str:
    return f"{d.year:04d}-{d.month:02d}"


def _meses(a: dt.date, b: dt.date) -> list[str]:
    out, y, m = [], a.year, a.month
    while (y, m) <= (b.year, b.month):
        out.append(f"{y:04d}-{m:02d}")
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def _distribuir(total: float, ini: dt.date, fim: dt.date) -> dict[str, float]:
    """Valor distribuído por mês, proporcional aos dias corridos entre início e término."""
    if fim < ini:
        ini, fim = fim, ini
    dias = (fim - ini).days + 1
    out: dict[str, float] = {}
    d = ini
    while d <= fim:
        prox = dt.date(d.year + (d.month == 12), d.month % 12 + 1, 1)
        ate = min(fim, prox - dt.timedelta(days=1))
        out[_mes(d)] = out.get(_mes(d), 0.0) + total * ((ate - d).days + 1) / dias
        d = prox
    return out


def _fracao(ini: dt.date, fim: dt.date, em: dt.date) -> float:
    if em < ini:
        return 0.0
    if em >= fim:
        return 1.0
    return ((em - ini).days + 1) / ((fim - ini).days + 1)


def _ua_das_tarefas(projeto) -> dict[str, list]:
    """Tarefas de cada código UA, lido do campo Texto "Código UA" do cronograma padrão Custo Plano
    (ex.: "GSE20", "GSE01/GSE03" ou "GSA" para todos os GSA). Cronograma sem esse campo: dicionário vazio."""
    from cronograma import Q, _filho, _txt
    ea = _filho(projeto.raiz, "ExtendedAttributes")
    fid = None
    for d in (ea.findall(Q + "ExtendedAttribute") if ea is not None else []):
        if norm(_txt(d, "Alias") or _txt(d, "FieldName")) == "codigo ua":
            fid = _txt(d, "FieldID")
    if not fid:
        return {}
    por_ua: dict[str, list] = {}
    for t in projeto.tarefas:
        if t.filhos:
            continue
        for x in t.el.findall(Q + "ExtendedAttribute"):
            if _txt(x, "FieldID") == fid:
                for cod in re.split(r"\s*[/,;]\s*|\s+e\s+", _txt(x, "Value")):
                    if cod.strip():
                        por_ua.setdefault(cod.strip().upper(), []).append(t)
    return por_ua


def _tarefas_da_ua(por_ua: dict, ua: str | None) -> list:
    """[(tarefa, fração do custo)] de um código UA. A fração segue o custo de cada tarefa no cronograma
    (linha de base, senão custo atual); sem custo, segue a duração."""
    if not ua or not por_ua:
        return []
    ua = ua.strip().upper()
    ts = por_ua.get(ua)
    if not ts:  # "GSA" no cronograma vale para GSA01, GSA02...
        fam = re.match(r"[A-Z]+", ua)
        ts = por_ua.get(fam[0]) if fam and fam[0] != ua else None
    if not ts:
        return []
    from cronograma import Q, _txt

    def peso(t):
        for b in t.el.findall(Q + "Baseline"):
            if _txt(b, "Number") in ("", "0") and _moeda(b, "Cost") > 0:
                return _moeda(b, "Cost")
        return _moeda(t.el, "Cost")
    ws = [peso(t) for t in ts]
    if sum(ws) <= 0:
        ws = [max(t.dur_h, 0.0) for t in ts]
    if sum(ws) <= 0:
        ws = [1.0] * len(ts)
    return [(t, w / sum(ws)) for t, w in zip(ts, ws)]


def analisar(dados: dict, lancamentos: list, projeto=None, atual: dict | None = None,
             data_status: dt.date | None = None) -> dict:
    atual = atual or {}
    ds = data_status or dt.date.today()
    bdi = float(dados.get("bdi") or 0)
    comps = dados.get("composicoes") or {}

    # ------------------------------------------------ tarefas do cronograma
    por_eap, por_nome, pct = {}, {}, {}
    if projeto is not None:
        A = projeto.agregar(atual, {}, dt.datetime.combine(ds, dt.time(17)))
        for t in projeto.tarefas:
            por_eap[t.eap] = t
            por_nome.setdefault(norm(t.nome), t)
            pct[t.uid] = projeto.pct(A[t.uid], "r") if t.filhos else float(atual.get(t.uid, 0))
    por_ua = _ua_das_tarefas(projeto) if projeto is not None else {}

    def datas_tarefa(t):
        i, f = t.lb_inicio or t.inicio, t.lb_termino or t.termino
        return (i.date() if i else None), (f.date() if f else None)

    # ------------------------------------------------------------ etapas
    titulos = [i for i in dados["itens"] if i["titulo"]]
    nivel_min = min((i["codigo"].count(".") for i in titulos if i["codigo"]), default=0)
    etapas, ordem = {}, []

    def etapa_de(item, ultimo):
        cod = item["codigo"]
        melhor = None
        for t in titulos:
            if t["codigo"] and cod.startswith(t["codigo"] + ".") and t["codigo"].count(".") == nivel_min:
                melhor = t
        return melhor or ultimo

    itens_out, ultimo_titulo, ultimo_top = [], None, None
    tot_tipo = {k: 0.0 for k in TIPOS}
    sem_vinculo = {"n": 0, "valor": 0.0}
    meses_prev: dict[str, float] = {}
    for it in dados["itens"]:
        if it["titulo"]:
            ultimo_titulo = it
            if not it["codigo"] or it["codigo"].count(".") == nivel_min:
                ultimo_top = it
            continue
        et = etapa_de(it, ultimo_top or ultimo_titulo)
        chave = (et["codigo"] + " " + et["descricao"]).strip() if et else "Sem etapa"
        if chave not in etapas:
            etapas[chave] = {"codigo": et["codigo"] if et else "", "nome": et["descricao"] if et else "Sem etapa",
                             "orcado": 0.0, "vp": 0.0, "va": 0.0, "cr": 0.0, "meses": {}, "titulo": et}
            ordem.append(chave)
        E = etapas[chave]

        comp = comps.get(it["comp"]) if it["comp"] else None
        por_tipo = {k: 0.0 for k in TIPOS}
        insumos = []
        if comp:
            for ins in comp["insumos"]:
                sub = ins["coef"] * ins["preco"]
                por_tipo[ins["tipo"]] += sub
                insumos.append(dict(ins, subtotal=round(sub, 4)))
            unit = sum(por_tipo.values())
        else:
            unit = float(it["preco"] or 0)
            if not unit and it.get("total_planilha") and it["qtd"]:
                unit = it["total_planilha"] / it["qtd"]
            por_tipo["out"] = unit
        direto = unit * float(it["qtd"] or 0)
        total = direto  # controle de custo: custo direto (o BDI entra só no preço de venda)
        for k in TIPOS:
            tot_tipo[k] += por_tipo[k] * float(it["qtd"] or 0)

        # vínculo com o cronograma: pelo código UA (orçamento e cronograma do padrão Custo Plano), pela EAP ou pelo nome.
        # Um grupo UA pode estar repartido em várias tarefas: o custo do item segue a mesma repartição do cronograma.
        partes = _tarefas_da_ua(por_ua, it.get("ua"))
        t = None
        if not partes:
            if it["tarefa"]:
                t = por_eap.get(it["tarefa"]) or por_nome.get(norm(it["tarefa"]))
            if t is None and it["codigo"] in por_eap:
                t = por_eap[it["codigo"]]
            if t is None and et is not None:
                if et.get("tarefa"):
                    t = por_eap.get(et["tarefa"]) or por_nome.get(norm(et["tarefa"]))
                t = t or por_eap.get(et["codigo"]) or por_nome.get(norm(et["descricao"]))
            if t is not None:
                partes = [(t, 1.0)]
        trechos = []  # (início, término, % físico, fração do custo)
        for tp, w in partes:
            a, b = datas_tarefa(tp)
            trechos.append((a, b, pct.get(tp.uid), w))
        if not trechos and it.get("inicio"):
            a = dt.date.fromisoformat(it["inicio"])
            trechos.append((a, dt.date.fromisoformat(it["termino"]) if it.get("termino") else a, None, 1.0))

        vp = va = 0.0
        for a, b, f, w in trechos:
            if a and b:
                for mes, v in _distribuir(total * w, a, b).items():
                    meses_prev[mes] = meses_prev.get(mes, 0.0) + v
                    E["meses"][mes] = E["meses"].get(mes, 0.0) + v
                vp += total * w * _fracao(a, b, ds)
            if f is not None:
                va += total * w * f / 100
        medidos = [(f, w) for _, _, f, w in trechos if f is not None]
        fisico = sum(f * w for f, w in medidos) / sum(w for _, w in medidos) if medidos and sum(w for _, w in medidos) else None
        inis = [a for a, _, _, _ in trechos if a]
        fins = [b for _, b, _, _ in trechos if b]
        ini, fim = (min(inis) if inis else None), (max(fins) if fins else None)
        if not partes and total:  # item sem custo não precisa de tarefa
            sem_vinculo["n"] += 1
            sem_vinculo["valor"] += total
        E["orcado"] += total
        E["vp"] += vp
        E["va"] += va
        rotulo = (partes[0][0].eap + " " + partes[0][0].nome) if partes else ""
        if len(partes) > 1:
            rotulo += f" (+{len(partes) - 1} tarefas)"
        itens_out.append({
            "codigo": it["codigo"], "descricao": it["descricao"], "unidade": it["unidade"], "qtd": it["qtd"],
            "comp": it["comp"], "comp_desc": comp["descricao"] if comp else "", "unit": unit, "direto": direto,
            "total": total, "venda": direto * (1 + bdi), "tipos": por_tipo, "insumos": insumos, "etapa": chave,
            "tarefa": rotulo, "fisico": fisico,
            "inicio": ini.isoformat() if ini else None, "termino": fim.isoformat() if fim else None, "vp": vp, "va": va,
        })

    # --------------------------------------------------------- realizado
    todos = [dict(x, origem="planilha") for x in dados.get("realizado") or []] + list(lancamentos or [])
    cod_item = {i["codigo"]: i for i in itens_out if i["codigo"]}
    meses_real: dict[str, float] = {}
    nao_ligados = 0.0
    cr_total = cr_ate = 0.0
    lanc_out = []
    for lc in todos:
        d = dt.date.fromisoformat(lc["data"])
        v = float(lc["valor"])
        cod = (lc.get("codigo") or "").strip()
        chave = None
        if cod in cod_item:
            chave = cod_item[cod]["etapa"]
        else:
            for k in ordem:
                ek = etapas[k]["codigo"]
                if cod and ek and (cod == ek or cod.startswith(ek + ".")):
                    chave = k
                    break
            if chave is None and cod:
                chave = next((k for k in ordem if norm(etapas[k]["nome"]) == norm(cod)), None)
        if chave:
            etapas[chave]["cr"] += v
        else:
            nao_ligados += v
        meses_real[_mes(d)] = meses_real.get(_mes(d), 0.0) + v
        cr_total += v
        if d <= ds:
            cr_ate += v
        lanc_out.append(dict(lc, etapa=chave or "Sem etapa"))
    lanc_out.sort(key=lambda x: x["data"], reverse=True)

    # ------------------------------------------------ série mensal
    todos_meses = sorted(set(meses_prev) | set(meses_real))
    serie = []
    if todos_meses:
        a = dt.date.fromisoformat(todos_meses[0] + "-01")
        b = dt.date.fromisoformat(todos_meses[-1] + "-01")
        acp = acr = 0.0
        for mes in _meses(a, b):
            p_, r_ = meses_prev.get(mes, 0.0), meses_real.get(mes, 0.0)
            acp += p_
            acr += r_
            serie.append({"mes": mes, "previsto": p_, "realizado": r_, "previsto_acum": acp, "realizado_acum": acr})

    bac = sum(E["orcado"] for E in etapas.values())
    vp = sum(E["vp"] for E in etapas.values())
    va = sum(E["va"] for E in etapas.values())
    cr = cr_total
    idc = va / cr if cr > 0 and (va or vp) else None  # sem nenhum avanço físico ligado ao custo, o IDC não existe
    idp = va / vp if vp > 0 else None
    eac = (cr + (bac - va) / idc) if idc else None
    etapas_out = []
    for k in ordem:
        E = etapas[k]
        etapas_out.append({"chave": k, "codigo": E["codigo"], "nome": E["nome"], "orcado": E["orcado"], "vp": E["vp"],
                           "va": E["va"], "cr": E["cr"], "idc": (E["va"] / E["cr"]) if E["cr"] > 0 and (va or vp) else None,
                           "meses": E["meses"]})
    if nao_ligados:
        etapas_out.append({"chave": "Sem etapa", "codigo": "", "nome": "Lançamentos sem etapa", "orcado": 0.0, "vp": 0.0,
                           "va": 0.0, "cr": nao_ligados, "idc": None, "meses": {}})
    direto = sum(tot_tipo.values())
    return {
        "arquivo": dados.get("arquivo"), "bdi": bdi, "data_status": ds.isoformat(), "tem_cronograma": projeto is not None,
        "totais": {"bac": bac, "venda": bac * (1 + bdi), "direto": direto, "vp": vp, "va": va, "cr": cr, "cr_ate": cr_ate, "idc": idc, "idp": idp,
                   "eac": eac, "vac": (bac - eac) if eac is not None else None},
        "tipos": [{"tipo": k, "nome": TIPOS[k], "valor": v, "com_bdi": v * (1 + bdi)} for k, v in tot_tipo.items() if v],
        "etapas": etapas_out, "itens": itens_out, "serie": serie, "lancamentos": lanc_out,
        "sem_vinculo": sem_vinculo, "avisos": dados.get("avisos") or [],
        # sem nenhum lançamento real: o realizado e o IDC não existem (não mostrar R$ 0 como se fosse gasto)
        "sem_realizado": not todos,
    }


# -------------------------------------------------- modelo e exemplo
CAB_ORC = ["Código", "Descrição", "Unidade", "Quantidade", "Composição", "Custo unitário", "Tarefa (EAP)", "Início", "Término"]
CAB_COMP = ["Composição", "Descrição da composição", "Unidade da composição", "Tipo", "Insumo", "Unidade do insumo",
            "Coeficiente", "Preço unitário"]
CAB_REAL = ["Data", "Código (item ou etapa)", "Descrição", "Fornecedor", "Tipo", "Valor"]


def montar_planilha(itens: list, comps: list, realizado: list, bdi: float, instrucoes: bool = True) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    azul, cinza = PatternFill("solid", fgColor="282B59"), PatternFill("solid", fgColor="E3EAE9")
    branco = Font(bold=True, color="FFFFFF")

    def aba(ws, cab, larg):
        ws.append(cab)
        for c in ws[1]:
            c.fill, c.font = azul, branco
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
        for i, w in enumerate(larg, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Orçamento"
    aba(ws, CAB_ORC, [10, 46, 9, 12, 13, 14, 13, 12, 12])
    for it in itens:
        ws.append(it)
        r = ws.max_row
        if it[3] in (None, "") and it[5] in (None, "") and it[4] in (None, ""):
            for c in ws[r]:
                c.font, c.fill = Font(bold=True), cinza
        for col in (8, 9):
            ws.cell(r, col).number_format = "dd/mm/yyyy"
        ws.cell(r, 6).number_format = '#,##0.00'

    wc = wb.create_sheet("Composições")
    aba(wc, CAB_COMP, [13, 40, 11, 14, 38, 10, 12, 13])
    dv = DataValidation(type="list", formula1='"Material,Mão de obra,Equipamento,Serviço,Outros"', allow_blank=True)
    wc.add_data_validation(dv)
    dv.add("D2:D5000")
    for c in comps:
        wc.append(c)
        wc.cell(wc.max_row, 8).number_format = '#,##0.00'

    wr = wb.create_sheet("Realizado")
    aba(wr, CAB_REAL, [12, 18, 40, 28, 14, 14])
    for x in realizado:
        wr.append(x)
        wr.cell(wr.max_row, 1).number_format = "dd/mm/yyyy"
        wr.cell(wr.max_row, 6).number_format = '#,##0.00'

    wp = wb.create_sheet("Parâmetros")
    wp.append(["BDI (%)", round(bdi * 100, 2)])
    wp["A1"].font = Font(bold=True)
    wp.column_dimensions["A"].width = 16
    if instrucoes:
        linhas = [
            "", "Como preencher",
            "Orçamento: uma linha por item. Linhas só com Código e Descrição (sem quantidade) são títulos de etapa (ex.: 1.2 Fundações).",
            "Composição: código da composição (aba Composições). Sem composição, preencha o Custo unitário.",
            "Tarefa (EAP): EAP da tarefa no MS Project (ex.: 1.3.2). Se ficar vazio, o sistema usa o código ou o nome da etapa.",
            "Início/Término: só se o item não estiver ligado ao cronograma.",
            "Composições: uma linha por insumo. Tipo = Material, Mão de obra, Equipamento, Serviço ou Outros.",
            "Custo unitário do item = soma de coeficiente × preço de cada insumo. O BDI é aplicado sobre o custo direto.",
            "Realizado (opcional): custos já pagos. Código = código do item ou da etapa. Também dá para lançar pelo sistema.",
        ]
        for t in linhas:
            wp.append([t])
        wp["A3"].font = Font(bold=True, size=12)
        wp.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def modelo() -> bytes:
    itens = [
        ["1", "SERVIÇOS PRELIMINARES", None, None, None, None, None, None, None],
        ["1.1", "Locação da obra", "m²", 420, "C-001", None, "1.1.2", None, None],
        ["1.2", "Placa de obra", "m²", 6, None, 310.0, None, dt.date(2026, 1, 5), dt.date(2026, 1, 9)],
        ["2", "ESTRUTURA", None, None, None, None, None, None, None],
        ["2.1", "Concreto armado fck 30 MPa", "m³", 60, "C-002", None, "1.3.1", None, None],
    ]
    comps = [
        ["C-001", "Locação convencional de obra", "m²", "Mão de obra", "Carpinteiro", "h", 0.06, 28.0],
        [None, None, None, "Mão de obra", "Servente", "h", 0.06, 22.5],
        [None, None, None, "Material", "Pontalete e tábua", "m", 0.35, 9.8],
        ["C-002", "Concreto armado fck 30 MPa (forma, armação e lançamento)", "m³", "Material", "Concreto usinado fck 30", "m³", 1.03, 540.0],
        [None, None, None, "Material", "Aço CA-50", "kg", 105, 8.9],
        [None, None, None, "Material", "Forma de compensado", "m²", 8, 62.0],
        [None, None, None, "Mão de obra", "Carpinteiro", "h", 7, 28.0],
        [None, None, None, "Mão de obra", "Armador", "h", 6, 28.0],
        [None, None, None, "Equipamento", "Bomba de concreto", "m³", 1, 45.0],
    ]
    real = [[dt.date(2026, 1, 20), "1.1", "Equipe de locação - medição 1", "Empreiteira Alfa", "Mão de obra", 1520.0]]
    return montar_planilha(itens, comps, real, 0.25)


COMP_EXEMPLO = [
    ("C-CANT", "Canteiro de obras (instalações provisórias)", "vb", [
        ("Serviço", "Locação de contêineres (mês)", "mês", 3, 2800.0), ("Mão de obra", "Servente", "h", 160, 22.5),
        ("Material", "Tapume de chapa compensada", "m²", 120, 48.0)]),
    ("C-LOC", "Locação convencional de obra", "m²", [
        ("Mão de obra", "Carpinteiro", "h", 0.06, 28.0), ("Mão de obra", "Servente", "h", 0.06, 22.5),
        ("Material", "Pontalete e tábua", "m", 0.35, 9.8)]),
    ("C-EST", "Estaca hélice contínua D40", "m", [
        ("Serviço", "Perfuração hélice contínua", "m", 1, 92.0), ("Material", "Concreto usinado fck 30", "m³", 0.14, 520.0),
        ("Material", "Aço CA-50", "kg", 3.2, 8.9)]),
    ("C-BLO", "Blocos e baldrames em concreto armado", "m³", [
        ("Material", "Concreto usinado fck 30", "m³", 1.05, 520.0), ("Material", "Aço CA-50", "kg", 90, 8.9),
        ("Material", "Forma de compensado", "m²", 6, 62.0), ("Mão de obra", "Pedreiro", "h", 6, 28.0),
        ("Mão de obra", "Servente", "h", 8, 22.5)]),
    ("C-ESTR", "Estrutura de concreto armado (pilares, vigas e laje)", "m³", [
        ("Material", "Concreto usinado fck 30", "m³", 1.03, 540.0), ("Material", "Aço CA-50", "kg", 105, 8.9),
        ("Material", "Forma de compensado", "m²", 8, 62.0), ("Mão de obra", "Carpinteiro", "h", 7, 28.0),
        ("Mão de obra", "Armador", "h", 6, 28.0), ("Mão de obra", "Servente", "h", 10, 22.5),
        ("Equipamento", "Bomba de concreto", "m³", 1, 45.0)]),
    ("C-ALV", "Alvenaria de bloco cerâmico 14 cm", "m²", [
        ("Material", "Bloco cerâmico 14x19x29", "un", 13.5, 2.35), ("Material", "Argamassa de assentamento", "m³", 0.012, 480.0),
        ("Mão de obra", "Pedreiro", "h", 0.75, 28.0), ("Mão de obra", "Servente", "h", 0.45, 22.5)]),
    ("C-HID", "Instalações hidrossanitárias (por pavimento)", "pav", [
        ("Material", "Tubos, conexões e registros", "vb", 1, 9800.0), ("Mão de obra", "Encanador", "h", 160, 30.0),
        ("Mão de obra", "Ajudante", "h", 160, 22.5)]),
    ("C-ELE", "Instalações elétricas (por pavimento)", "pav", [
        ("Material", "Cabos, eletrodutos e caixas", "vb", 1, 11500.0), ("Mão de obra", "Eletricista", "h", 180, 30.0),
        ("Mão de obra", "Ajudante", "h", 180, 22.5)]),
    ("C-LIMP", "Limpeza final de obra", "m²", [
        ("Mão de obra", "Servente", "h", 0.3, 22.5), ("Material", "Material de limpeza", "m²", 1, 1.8)]),
]


def exemplo(projeto=None, atual: dict | None = None) -> bytes:
    """Orçamento de exemplo ligado ao cronograma de exemplo (Residencial Exemplo – Torre A)."""
    itens = [["1.1", "Serviços preliminares", None, None, None, None, None, None, None],
             ["1.1.1", "Canteiro de obras", "vb", 1, "C-CANT", None, "1.1.1", None, None],
             ["1.1.2", "Locação da obra", "m²", 420, "C-LOC", None, "1.1.2", None, None],
             ["1.2", "Fundações", None, None, None, None, None, None, None],
             ["1.2.1", "Estacas hélice contínua D40", "m", 640, "C-EST", None, "1.2.1", None, None],
             ["1.2.2", "Blocos e baldrames", "m³", 38, "C-BLO", None, "1.2.2", None, None],
             ["1.3", "Estrutura", None, None, None, None, None, None, None]]
    itens += [[f"1.3.{n}", f"Estrutura de concreto armado - Pav. {n}", "m³", 62, "C-ESTR", None, f"1.3.{n}", None, None] for n in range(1, 6)]
    itens += [["1.4", "Alvenaria", None, None, None, None, None, None, None]]
    itens += [[f"1.4.{n}", f"Alvenaria de vedação - Pav. {n}", "m²", 380, "C-ALV", None, f"1.4.{n}", None, None] for n in range(1, 6)]
    itens += [["1.5", "Instalações", None, None, None, None, None, None, None]]
    for n in range(1, 6):
        itens += [[f"1.5.{n}.1", f"Instalações hidrossanitárias - Pav. {n}", "pav", 1, "C-HID", None, f"1.5.{n}", None, None],
                  [f"1.5.{n}.2", f"Instalações elétricas - Pav. {n}", "pav", 1, "C-ELE", None, f"1.5.{n}", None, None]]
    itens += [["1.6", "Entrega", None, None, None, None, None, None, None],
              ["1.6.1", "Limpeza final", "m²", 2100, "C-LIMP", None, "1.6", None, None]]
    comps = []
    for cod, desc, un, ins in COMP_EXEMPLO:
        for k, (tipo, nome, uni, coef, preco) in enumerate(ins):
            comps.append([cod if k == 0 else None, desc if k == 0 else None, un if k == 0 else None, tipo, nome, uni, coef, preco])

    # realizado de exemplo: proporcional ao % físico de cada tarefa, com desvios típicos
    real = []
    if projeto is not None:
        dados = ler_custos("exemplo.xlsx", montar_planilha(itens, comps, [], 0.25, instrucoes=False))
        an = analisar(dados, [], projeto, atual or projeto.medicao_arquivo, dt.date.today())
        fator = {"1.1": 1.04, "1.2": 0.97, "1.3": 1.09, "1.4": 1.02, "1.5": 1.12}
        fornec = {"1.1": "Equipe própria", "1.2": "Fundações Solo Firme", "1.3": "Empreiteira Beta",
                  "1.4": "Empreiteira Alfa", "1.5": "Instaladora Gama"}
        for it in an["itens"]:
            if not it["fisico"] or not it["termino"]:
                continue
            et = ".".join(it["codigo"].split(".")[:2])
            v = it["direto"] * it["fisico"] / 100 * fator.get(et, 1.0)
            dia = min(dt.date.fromisoformat(it["termino"]), dt.date.today())
            mat = v * (it["tipos"]["mat"] / it["unit"]) if it["unit"] else 0
            if mat > 1:
                real.append([dia - dt.timedelta(days=5), it["codigo"], "Materiais - " + it["descricao"], "Fornecedores diversos", "Material", round(mat, 2)])
            if v - mat > 1:
                real.append([dia, it["codigo"], "Medição de mão de obra - " + it["descricao"], fornec.get(et, "Empreiteira"), "Serviço", round(v - mat, 2)])
    return montar_planilha(itens, comps, real, 0.25)


MESES_BR = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]


def rotulo_mes(m: str) -> str:
    return f"{MESES_BR[int(m[5:7]) - 1]}/{m[2:4]}"


def exportar_xlsx(a: dict) -> bytes:
    """Cronograma financeiro (etapas × meses) e orçado × realizado em Excel."""
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    azul, cinza = PatternFill("solid", fgColor="282B59"), PatternFill("solid", fgColor="E3EAE9")
    branco, neg = Font(bold=True, color="FFFFFF"), Font(bold=True)
    moeda = '#,##0.00'
    meses = [s["mes"] for s in a["serie"]]

    ws = wb.active
    ws.title = "Cronograma financeiro"
    ws.append(["Cronograma financeiro (custo direto)", None, f"Data de status {a['data_status'][8:10]}/{a['data_status'][5:7]}/{a['data_status'][:4]}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Código", "Etapa"] + [rotulo_mes(m) for m in meses] + ["Total"])
    for c in ws[3]:
        c.fill, c.font = azul, branco
        c.alignment = Alignment(horizontal="center")
    for e in a["etapas"]:
        if not e["orcado"]:
            continue
        ws.append([e["codigo"], e["nome"]] + [round(e["meses"].get(m, 0.0), 2) for m in meses] + [round(e["orcado"], 2)])
    lin = ["", "Previsto no mês"] + [round(s["previsto"], 2) for s in a["serie"]] + [round(a["totais"]["bac"], 2)]
    ws.append(lin)
    ws.append(["", "Previsto acumulado"] + [round(s["previsto_acum"], 2) for s in a["serie"]])
    ws.append(["", "Realizado no mês"] + [round(s["realizado"], 2) for s in a["serie"]] + [round(a["totais"]["cr"], 2)])
    ws.append(["", "Realizado acumulado"] + [round(s["realizado_acum"], 2) for s in a["serie"]])
    for r in range(ws.max_row - 3, ws.max_row + 1):
        for c in ws[r]:
            c.fill, c.font = cinza, neg
    for row in ws.iter_rows(min_row=4, min_col=3):
        for c in row:
            c.number_format = moeda
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 34
    for i in range(3, len(meses) + 4):
        ws.column_dimensions[ws.cell(3, i).column_letter].width = 14
    ws.freeze_panes = "C4"

    w2 = wb.create_sheet("Orçado x realizado")
    w2.append(["Código", "Etapa", "Orçado", "Previsto até a data", "Valor agregado", "Realizado", "Variação de custo (VA − CR)", "IDC"])
    for c in w2[1]:
        c.fill, c.font = azul, branco
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for e in a["etapas"]:
        w2.append([e["codigo"], e["nome"], e["orcado"], e["vp"], e["va"], e["cr"], e["va"] - e["cr"], e["idc"]])
    t = a["totais"]
    w2.append(["", "Total", t["bac"], t["vp"], t["va"], t["cr"], t["va"] - t["cr"], t["idc"]])
    for c in w2[w2.max_row]:
        c.fill, c.font = cinza, neg
    for row in w2.iter_rows(min_row=2, min_col=3, max_col=7):
        for c in row:
            c.number_format = moeda
    for row in w2.iter_rows(min_row=2, min_col=8, max_col=8):
        for c in row:
            c.number_format = "0.00"
    for col, w in zip("ABCDEFGH", (9, 34, 15, 16, 15, 15, 18, 8)):
        w2.column_dimensions[col].width = w
    w2.row_dimensions[1].height = 30

    w3 = wb.create_sheet("Lançamentos")
    w3.append(["Data", "Código", "Etapa", "Descrição", "Fornecedor", "Tipo", "Valor", "Origem"])
    for c in w3[1]:
        c.fill, c.font = azul, branco
    for x in a["lancamentos"]:
        w3.append([dt.date.fromisoformat(x["data"]), x.get("codigo"), x.get("etapa"), x.get("descricao"), x.get("fornecedor"),
                   TIPOS.get(x.get("tipo"), "Outros"), x["valor"], x.get("origem")])
        w3.cell(w3.max_row, 1).number_format = "dd/mm/yyyy"
        w3.cell(w3.max_row, 7).number_format = moeda
    for col, w in zip("ABCDEFGH", (11, 10, 26, 44, 26, 13, 14, 10)):
        w3.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------- orçamento a partir do cronograma
def _moeda(el, nome) -> float:
    """Valores de custo no XML do MS Project vêm em centavos."""
    from cronograma import _txt
    v = numero(_txt(el, nome))
    return (v or 0.0) / 100


def dados_do_projeto(p) -> dict | None:
    """Orçamento tirado das colunas de custo do próprio cronograma (Custo e Custo da linha de base).
    O Custo real do Project não entra como realizado: o Project calcula esse valor sozinho (% concluído × custo),
    e ele não é gasto real. O realizado vem só dos lançamentos. Devolve None se o cronograma não tiver custos."""
    from cronograma import Q, _txt
    raizes = [t for t in p.tarefas if t.pai is None]
    pular = {raizes[0].uid} if len(raizes) == 1 and raizes[0].filhos else set()
    itens, custo_real_project = [], 0.0
    for t in p.tarefas:
        if t.uid in pular:
            continue
        if t.filhos:
            itens.append({"codigo": t.eap, "descricao": t.nome, "titulo": True, "unidade": "", "qtd": 0.0,
                          "comp": "", "preco": None, "total_planilha": None, "tarefa": t.eap,
                          "inicio": None, "termino": None})
            continue
        custo = _moeda(t.el, "Cost")
        lb = 0.0
        for b in t.el.findall(Q + "Baseline"):
            if _txt(b, "Number") in ("", "0"):
                lb = _moeda(b, "Cost")
        orcado = lb or custo
        if orcado > 0:
            itens.append({"codigo": t.eap, "descricao": t.nome, "titulo": False, "unidade": "vb", "qtd": 1.0,
                          "comp": "", "preco": round(orcado, 2), "total_planilha": round(orcado, 2), "tarefa": t.eap,
                          "inicio": None, "termino": None})
        custo_real_project += _moeda(t.el, "ActualCost")
    if not any(not i["titulo"] for i in itens):
        return None
    avisos = []
    if custo_real_project > 0:
        avisos.append("O cronograma traz a coluna Custo real do Project, mas ela não foi usada como realizado: o Project "
                      "calcula esse valor sozinho (% concluído × custo) e ele não é gasto real. O realizado vem só dos "
                      "custos lançados aqui.")
    return {"arquivo": p.arquivo, "origem": "project", "bdi": 0.0, "itens": itens, "composicoes": {},
            "realizado": [], "avisos": avisos}


def _zip_seguro(conteudo: bytes, limite: int = 150_000_000) -> bool:
    """Uma planilha .xlsx é um zip: recusa as que, descompactadas, passariam do limite (zip bomb)."""
    import zipfile
    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
            return sum(i.file_size for i in z.infolist()) <= limite
    except zipfile.BadZipFile:
        return True  # o openpyxl dá o erro de arquivo inválido


# ------------------------------------------------------- orçado x realizado
# Estrutura da planilha de controle da Engenharia (projetos/_modelo/orcamento/orcado_realizado.py):
# (a) orçamento inicial, (b) atualizado, (c) contratado, (d) incorrido, (e) saldo de contrato, (f) saldo de pedido,
# (h) saldo a contratar (premissa). O sistema calcula (g) = d+e+f, (i) = g+h, (j) = b−i e (k) = i÷b.
CONTROLE_ENTRADAS = {"a": "oin", "b": "oatu", "c": "contratado", "d": "incorrido", "e": "saldo_contrato",
                     "f": "saldo_pedido", "h": "a_contratar"}


def _linhas_csv(conteudo: bytes) -> list:
    import csv
    if len(conteudo) > 5_000_000:
        raise ErroCustos("O arquivo é grande demais para ser lido.")
    for cod in ("utf-8-sig", "cp1252"):
        try:
            texto_ = conteudo.decode(cod)
            break
        except UnicodeDecodeError:
            continue
    amostra = texto_[:4000]
    sep = ";" if amostra.count(";") >= amostra.count(",") else ","
    if "\t" in amostra and amostra.count("\t") > amostra.count(sep):
        sep = "\t"
    return [r for r in csv.reader(io.StringIO(texto_), delimiter=sep)][:6000]


def ler_controle(nome_arquivo: str, conteudo: bytes) -> dict:
    """Aceita Excel (.xlsx, .xlsm; aba "Orçado x Realizado" ou a única aba) ou texto separado (.csv, .txt)."""
    nome = nome_arquivo or ""
    if re.search(r"\.(csv|txt)$", nome, re.I):
        L, extra = _linhas_csv(conteudo), None
    elif re.search(r"\.xls[xm]$", nome, re.I):
        try:
            if not _zip_seguro(conteudo):
                raise ErroCustos("A planilha é grande demais para ser lida.")
            wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
        except Exception:
            raise ErroCustos("Não consegui abrir a planilha. Confira se é um arquivo .xlsx válido.")
        try:
            ws = next((w for w in wb.worksheets if "orcado x realizado" in norm(w.title).replace("×", "x")), None)
            if ws is None and len(wb.worksheets) == 1:
                ws = wb.worksheets[0]
            if ws is None:
                raise ErroCustos("Não achei a aba Orçado x Realizado. Use a planilha de controle no padrão Custo Plano.")
            L = _linhas(ws, 6000, 20)
            extra = _erp_orcamento_analitico(wb)
        finally:
            wb.close()
    elif re.search(r"\.xls$", nome, re.I):
        raise ErroCustos("Arquivos .xls (Excel antigo) não são lidos. No Excel, use Salvar como > Pasta de Trabalho do Excel (.xlsx) ou CSV.")
    else:
        raise ErroCustos("Envie a planilha Orçado x Realizado em Excel (.xlsx, .xlsm) ou CSV (.csv).")
    erp = _ler_controle_erp(L)
    if erp:
        if extra:
            extra["arquivo"] = nome_arquivo
            erp["orcamento"] = extra
        return dict(erp, arquivo=nome_arquivo)
    r0 = col = None
    for r, linha in enumerate(L[:40]):
        m = {}
        for c, v in enumerate(linha):
            n = norm(v)
            achou = re.search(r"\(([a-k])\)\s*$", n)
            if achou:
                m[achou[1]] = c
            elif n.startswith("cod"):
                m["cod"] = c
            elif n.startswith("servico") or n.startswith("descricao"):
                m["serv"] = c
        if all(k in m for k in ("serv", "a", "b", "c", "d", "e", "f", "h")):
            r0, col = r, m
            break
    if r0 is None:
        raise ErroCustos("Na aba Orçado x Realizado não encontrei o cabeçalho com Serviço e as colunas (a) a (h).")
    linhas, avisos, vazias = [], [], []
    for linha in L[r0 + 1:]:
        serv = texto(linha[col["serv"]]) if col["serv"] < len(linha) else ""
        cod = texto(linha[col["cod"]]) if "cod" in col and col["cod"] < len(linha) else ""
        if norm(serv).startswith("total") or norm(cod).startswith("total"):
            break
        if not serv and not cod:
            if linhas:
                break
            continue
        x = {"codigo": cod, "servico": serv}
        for letra, campo in CONTROLE_ENTRADAS.items():
            v = linha[col[letra]] if col[letra] < len(linha) else None
            if isinstance(v, str) and v.strip() in ("-", "–", "—"):
                v = 0
            n = numero(v)
            if n is None and v not in (None, ""):
                vazias.append(f"{cod or serv} ({letra})")
            x[campo] = int(_arred(n or 0))
        linhas.append(x)
    if not linhas:
        raise ErroCustos("A aba Orçado x Realizado não tem serviços preenchidos.")
    if vazias:
        avisos.append("Valores que não consegui ler e ficaram zerados: " + ", ".join(vazias[:8]) + ("…" if len(vazias) > 8 else ""))
    return {"arquivo": nome_arquivo, "linhas": linhas, "avisos": avisos}


ERP_COLS = {"oin": ("orcado inicial",), "contratado": ("contratado",), "saldo_pedido": ("saldo de pedido",),
            "saldo_contrato": ("saldo de contrato",), "ajustes": ("ajustes",), "realizado": ("realizado",),
            "a_pagar": ("a pagar",), "apropriado": ("apropriado financeiro",), "tendencia": ("orcado tendencia",)}


def _ler_controle_erp(L) -> dict | None:
    """Orçado x Realizado exportado do ERP da construtora (colunas Ref., Descrição / Credor, Orçado Inicial,
    Contratado, Saldo de Pedido, Saldo de Contrato, Ajustes, Orçado Tendência, Saldo de Orçamento, Realizado,
    A Pagar, Apropriado financeiro). Lê uma linha por etapa de cada unidade de custo (UC) e leva para (a)...(h):
    (b) = orçado inicial + ajustes; (d) = apropriado financeiro (realizado + a pagar); (h) não existe no ERP e
    entra como premissa (regra do gerente de orçamento, 25/10/2026): zero com contratado ≥ 95% de (b); (b) − (c)
    quando falta contratar parte do escopo; (b) − (g) sem contrato. A tendência do ERP não é usada: ela repete o contratado."""
    r0 = m = None
    for r, linha in enumerate(L[:40]):
        nomes = [norm(v) for v in linha]
        if not any(n.startswith("ref") for n in nomes) or not any(n.startswith("descricao") for n in nomes):
            continue
        mm = {k: nomes.index(n) for k, ops in ERP_COLS.items() for n in nomes if n in ops}
        if all(k in mm for k in ("oin", "contratado", "saldo_pedido", "saldo_contrato", "realizado")):
            r0, m = r, mm
            m["ref"] = next(i for i, n in enumerate(nomes) if n.startswith("ref"))
            m["desc"] = next(i for i, n in enumerate(nomes) if n.startswith("descricao"))
            break
    if r0 is None:
        return None
    val = lambda linha, k: int(_arred(numero(linha[m[k]]) or 0)) if k in m and m[k] < len(linha) else 0  # noqa: E731
    def entradas(linha, com_h=True):
        a, c = val(linha, "oin"), val(linha, "contratado")
        b = a + val(linha, "ajustes")
        d = val(linha, "apropriado") if "apropriado" in m else val(linha, "realizado") + val(linha, "a_pagar")
        e, f = val(linha, "saldo_contrato"), val(linha, "saldo_pedido")
        g = d + e + f
        if not com_h:  # credor: não tem orçamento próprio, nada a contratar
            h = 0
        elif c and c >= 0.95 * b:  # escopo considerado todo contratado
            h = 0
        elif c:  # falta contratar parte do escopo: valorizada pelo orçamento
            h = max(0, b - c)
        else:  # sem contrato
            h = max(0, b - g)
        return {"oin": a, "oatu": b, "contratado": c, "incorrido": d, "saldo_contrato": e, "saldo_pedido": f, "a_contratar": h}

    # níveis do ERP: UC > etapa (Ref. de 2 dígitos) > itens (Ref. 01.000.000.001, aninhados pelo prefixo) > credor
    # (contrato CTPS/…, Pedido …, ou pagamento sem documento). As linhas por etapa continuam sendo a base de (a)…(k).
    linhas, niveis, uc, total_uc, pilha, no_uc = [], [], None, 0, [], None
    for linha in L[r0 + 1:]:
        ref, desc = texto(linha[m["ref"]]), texto(linha[m["desc"]])
        if not ref and not desc:
            continue
        uc_m = re.match(r"(\d+)\s*-\s*(.+)", desc) if not ref else None
        if uc_m and numero(linha[m["oin"]]) is not None:
            uc, pilha = uc_m[1], []
            total_uc += val(linha, "oin")
            no_uc = dict(entradas(linha), nivel=0, tipo="uc", codigo=uc, descricao=uc_m[2].strip(), a_contratar=0)
            niveis.append(no_uc)
            continue
        if uc is None:
            continue
        if re.fullmatch(r"\d{1,3}", ref):
            x = entradas(linha)
            linhas.append(dict(x, codigo=f"{uc}.{ref}", servico=desc))
            no_uc["a_contratar"] += x["a_contratar"]
            pilha = [ref]
            niveis.append(dict(x, nivel=1, tipo="etapa", codigo=f"{uc}.{ref}", descricao=desc))
        elif re.fullmatch(r"\d+(\.\d+)+", ref) and pilha:
            while len(pilha) > 1 and not ref.startswith(pilha[-1] + "."):
                pilha.pop()
            pilha.append(ref)
            niveis.append(dict(entradas(linha), nivel=len(pilha), tipo="item", codigo=f"{uc}.{ref}", descricao=desc))
        elif pilha:
            niveis.append(dict(entradas(linha, False), nivel=len(pilha) + 1, tipo="credor", codigo=ref, descricao=desc))
    if not linhas:
        return None
    avisos = ["Planilha do ERP: o incorrido (d) é o apropriado financeiro (realizado + a pagar). O saldo a contratar (h) "
              "não vem do ERP (a tendência do ERP repete o contratado) e é premissa: zero quando o contratado já passa de 95% de (b); "
              "(b) − (c) quando falta contratar parte do escopo; (b) − (g) quando não há contrato."]
    dif = [x["incorrido"] + x["saldo_contrato"] + x["saldo_pedido"] - x["contratado"] for x in linhas if x["contratado"]]
    acima = [v for v in dif if v > 1]
    if acima:  # bruto (só as etapas acima do contratado) e líquido da obra, para ninguém comparar um com o outro
        avisos.append(f"R$ {sum(acima):,.0f} acima do contratado em {len(acima)} etapas (líquido R$ {sum(dif):,.0f} na obra): "
                      "o incorrido mais os saldos passa do contratado (confira a aba Fora do Orçamento do ERP).".replace(",", "."))
    if abs(sum(x["oin"] for x in linhas) - total_uc) > len(linhas):
        avisos.append("A soma das etapas não fecha com o total das unidades de custo da planilha. Confira a exportação do ERP.")
    return {"linhas": linhas, "niveis": niveis, "avisos": avisos, "origem": "erp"}


def _erp_orcamento_analitico(wb) -> dict | None:
    """Abas Orçamento e Analítico do arquivo do ERP viram o orçamento da obra (itens por etapa) e os custos realizados
    com data, no mesmo formato do orçamento importado. Assim a tela Custos volta a ter Orçamento, Financeiro mês a mês e
    Valor agregado sem o Project. Realizado mês a mês = pago, pela data da baixa. "A pagar" não tem data: entra no
    incorrido (d), não no mês a mês. "Previsão (não entra)" fica de fora, como no ERP."""
    abas = {norm(w.title): w for w in wb.worksheets}
    if "orcamento" not in abas:
        return None
    cod_uc = lambda v: (re.match(r"\s*(\d+)\s*-", texto(v)) or [None, None])[1]  # noqa: E731
    L = _linhas(abas["orcamento"], 20000, 12)
    r0 = next((r for r, x in enumerate(L[:30]) if [norm(v) for v in x[:5]] == ["uc", "etapa", "subetapa", "apropriacao", "descricao"]), None)
    if r0 is None:
        return None
    itens, vistos = [], set()

    def titulo(cod, desc):
        if cod not in vistos:
            vistos.add(cod)
            itens.append({"codigo": cod, "descricao": desc, "titulo": True, "unidade": "", "qtd": 0.0, "comp": "",
                          "preco": None, "total_planilha": None, "tarefa": "", "inicio": None, "termino": None})
    for x in L[r0 + 1:]:
        x = list(x) + [None] * 9
        uc, ap = cod_uc(x[0]), re.sub(r"\D", "", texto(x[3]).split(".")[0])
        if not uc or len(ap) < 10:
            continue
        et = re.match(r"\s*(\d+)\s*-\s*(.*)", texto(x[1]))
        if not et:
            continue
        cet = f"{uc}.{int(et[1]):02d}"
        titulo(cet, et[2].strip())
        sub = re.match(r"\s*\d+\.(\d+)\s*-\s*(.*)", texto(x[2]))
        if sub:
            titulo(f"{cet}.{sub[1]}", sub[2].strip())
        valor = numero(x[8]) or 0.0
        qtd = numero(x[6]) or 0.0
        itens.append({"codigo": f"{cet}.{ap[-9:-6]}.{ap[-6:-3]}.{ap[-3:]}", "descricao": texto(x[4]), "titulo": False,
                      "unidade": texto(x[5]), "qtd": qtd or (1.0 if valor else 0.0), "comp": "", "preco": valor / qtd if qtd else valor,
                      "total_planilha": round(valor, 2), "tarefa": "", "inicio": None, "termino": None})
    if not any(not i["titulo"] for i in itens):
        return None
    realizado, a_pagar = [], 0.0
    if "analitico" in abas:
        A = _linhas(abas["analitico"], 40000, 10)
        a0 = next((r for r, x in enumerate(A[:30]) if any(norm(v).startswith("data da baixa") for v in x)), None)
        if a0 is not None:
            m = {norm(v): c for c, v in enumerate(A[a0])}
            ci = lambda *ks: next((m[k] for k in m for p in ks if k.startswith(p)), None)  # noqa: E731
            c_sit, c_uc, c_serv, c_doc, c_cred, c_data, c_val = (ci("situacao"), ci("uc"), ci("servico"), ci("documento"),
                                                                 ci("credor"), ci("data da baixa"), ci("considerado"))
            for x in A[a0 + 1:]:
                x = list(x) + [None] * 10
                sit = norm(x[c_sit])
                v = numero(x[c_val]) or 0.0
                uc = cod_uc(x[c_uc])
                if not uc or not v or sit.startswith("previsao") or sit == "total":
                    continue
                d = data(x[c_data])
                if d is None:  # a pagar: sem data da baixa
                    a_pagar += v
                    continue
                serv = texto(x[c_serv]).split(" - ")[0].strip()
                realizado.append({"data": d.isoformat(), "codigo": f"{uc}.{serv}" if re.fullmatch(r"[\d.]+", serv) else uc,
                                  "descricao": texto(x[c_doc]), "fornecedor": texto(x[c_cred])[:120], "tipo": "out",
                                  "valor": round(v, 2)})
    avisos = ["Orçamento e custos tirados do arquivo do ERP: o orçamento vem da aba Orçamento e o realizado mês a mês da aba "
              "Analítico (o que foi pago, pela data da baixa)."]
    if a_pagar:
        avisos.append(f"R$ {a_pagar:,.0f}".replace(",", ".") + " a pagar ainda não têm data da baixa: entram no incorrido (d) "
                      "do Orçado x Realizado, mas não no mês a mês.")
    return {"origem": "erp", "bdi": 0.0, "itens": itens, "composicoes": {}, "realizado": realizado,
            "a_pagar": round(a_pagar, 2), "avisos": avisos}


def calcular_controle(dados: dict | None, orcado_sistema: float | None = None, realizado_sistema: float | None = None):
    """Linhas com (g), (i), (j) e (k) calculados, o total da obra e os avisos de conferência com o sistema."""
    if not dados or not dados.get("linhas"):
        return None
    campos = list(CONTROLE_ENTRADAS.values()) + ["comprometido", "estimativa", "desvio"]
    out, tot = [], dict.fromkeys(campos, 0)
    for x in dados["linhas"]:
        y = _calc_linha(x)
        for k in campos:
            tot[k] += y[k]
        out.append(y)
    tot["indice"] = round(tot["estimativa"] / tot["oatu"], 3) if tot["oatu"] else None
    avisos = list(dados.get("avisos") or [])
    br = lambda v: f"R$ {v:,.0f}".replace(",", ".")  # noqa: E731
    if orcado_sistema and abs(tot["oin"] - orcado_sistema) > len(out):
        avisos.append(f"O orçamento inicial da planilha ({br(tot['oin'])}) é diferente do orçamento do sistema "
                      f"({br(orcado_sistema)}). Confira se as duas estão na mesma revisão.")
    if realizado_sistema and abs(tot["incorrido"] - realizado_sistema) > len(out):
        avisos.append(f"O incorrido da planilha ({br(tot['incorrido'])}) é diferente do custo realizado lançado no "
                      f"sistema ({br(realizado_sistema)}). Confira os lançamentos ou a planilha.")
    return {"arquivo": dados.get("arquivo"), "atualizado_em": dados.get("atualizado_em"), "linhas": out,
            "total": tot, "avisos": avisos, "tem_niveis": bool(dados.get("niveis"))}


def _calc_linha(x: dict) -> dict:
    y = dict(x)
    y["comprometido"] = x["incorrido"] + x["saldo_contrato"] + x["saldo_pedido"]
    y["estimativa"] = y["comprometido"] + x["a_contratar"]
    y["desvio"] = x["oatu"] - y["estimativa"]
    y["indice"] = round(y["estimativa"] / x["oatu"], 3) if x["oatu"] else None
    return y


def calcular_niveis(dados: dict | None) -> list | None:
    """Todos os níveis do ERP (UC, etapa, item, credor). (h) a (k) só em UC e etapa, onde a premissa é calculada;
    item e credor mostram os valores do ERP (a) a (g)."""
    if not dados or not dados.get("niveis"):
        return None
    out = [_calc_linha(x) for x in dados["niveis"]]
    for y in out:
        if y["tipo"] in ("item", "credor"):  # a premissa (h) e a estimativa são feitas por etapa (regra validada)
            y["a_contratar"] = y["estimativa"] = y["desvio"] = y["indice"] = None
    return out
