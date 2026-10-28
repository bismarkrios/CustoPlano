"""Teste da calculadora de R$/m²: conferência do quadro de áreas, R$/m², Excel e acesso só da equipe.
Usa o quadro da prancha que a Engenharia mandou (3.542,67 m², sem área privativa).
Rodar com:  CP_BANCO=/tmp/areas.db CP_SECRET=x CP_COPIAS_AUTO=0 python teste_areas.py
"""
import copy
import io
import os
import sys

os.environ.pop("ANTHROPIC_API_KEY", None)
import app
import areas
import banco

falhas = []


def checar(cond, msg):
    if not cond:
        falhas.append(msg)


def pav(nome, c, nc, t):
    return {"nome": nome, "computavel": c, "nao_computavel": nc, "total": t, "privativa": None, "unidades": None}


QUADRO = {
    "encontrado": True,
    "pavimentos": [pav("Sub-solo", None, 567.60, 567.60), pav("Térreo", 29.74, 232.91, 262.65),
                   pav("Tipo (1º ao 10º)", 2381.70, 157.40, 2539.10), pav("11º", 72.32, 24.34, 96.66),
                   pav("Barrilete", None, 38.33, 38.33), pav("Caixa d'água", None, 38.33, 38.33)],
    "total_computavel": 2483.76, "total_nao_computavel": 1058.91, "total_geral": 3542.67,
    "privativa_total": None, "unidades_total": None, "terreno": 690.0, "ca": 3.60, "to": 0.38, "observacoes": "",
}

# conferência
v = areas.validar(QUADRO)
checar(v["ok"], f"o quadro da prancha não fechou: {v}")
checar(v["ca"] and v["ca"]["ok"], "o CA de 3,60 não conferiu com computável ÷ terreno")
ruim = copy.deepcopy(QUADRO)
ruim["pavimentos"][2]["total"] = 2593.10  # erro de leitura (9 no lugar de 3)
v2 = areas.validar(ruim)
checar(not v2["ok"] and not v2["linhas"][2]["ok"] and all(l["ok"] for i, l in enumerate(v2["linhas"]) if i != 2), "não marcou só a linha errada")
checar(not v2["totais"]["total"]["ok"], "a soma com a linha errada bateu com o total impresso")
falta = copy.deepcopy(QUADRO)
falta["pavimentos"].pop()  # caixa d'água esquecida
checar(not areas.validar(falta)["ok"], "aceitou um quadro com linha faltando")

# R$/m²: construído sim, privativo não (não consta no projeto)
r = areas.calcular(QUADRO, 35_426_700, 17_713_350)
checar(round(r["venda"]["construido"]) == 10000 and round(r["custo"]["construido"]) == 5000, f"R$/m² construído errado: {r}")
checar(r["privativa"] is None and r["venda"]["privativo"] is None, "estimou área privativa que não consta no projeto")
com_priv = copy.deepcopy(QUADRO)
com_priv["privativa_total"] = 2000.0
checar(round(areas.calcular(com_priv, 30_000_000, None)["venda"]["privativo"]) == 15000, "R$/m² privativo errado")

# sem a chave, a leitura avisa que está desligada
banco.criar_usuario("equipe@teste.com", "senha-forte-123", "Equipe", admin=True)
banco.criar_usuario("cliente@teste.com", "senha-forte-123", "Cliente")
adm, cli = app.app.test_client(), app.app.test_client()
assert adm.post("/api/login", json={"email": "equipe@teste.com", "senha": "senha-forte-123"}).status_code == 200
assert cli.post("/api/login", json={"email": "cliente@teste.com", "senha": "senha-forte-123"}).status_code == 200
checar(adm.get("/api/admin/areas").get_json() == {"leitura": False}, "disse que a leitura está ligada sem chave")
pdf = open(os.path.join(app.AQUI, "web", "apresentacao.pdf"), "rb").read()[:2000]
r = adm.post("/api/admin/areas/ler", data={"arquivo": (io.BytesIO(pdf), "prancha.pdf", "application/pdf")}, content_type="multipart/form-data")
checar(r.status_code == 503 and "desligada" in r.get_json()["erro"], f"sem chave não avisou: {r.status_code} {r.get_json()}")

# com a leitura simulada, o quadro volta junto com a conferência
os.environ["ANTHROPIC_API_KEY"] = "teste"
original = areas.ler_quadro
areas.ler_quadro = lambda conteudo, mime: copy.deepcopy(QUADRO)
r = adm.post("/api/admin/areas/ler", data={"arquivo": (io.BytesIO(pdf), "prancha.pdf", "application/pdf")}, content_type="multipart/form-data")
checar(r.status_code == 200 and r.get_json()["conferencia"]["ok"] and len(r.get_json()["quadro"]["pavimentos"]) == 6, "a leitura simulada não voltou")
areas.ler_quadro = original
r = adm.post("/api/admin/areas/ler", data={"arquivo": (io.BytesIO(b"oi"), "planilha.xlsx", "application/vnd.ms-excel")}, content_type="multipart/form-data")
checar(r.status_code == 400, "aceitou um arquivo que não é PDF nem imagem")
os.environ.pop("ANTHROPIC_API_KEY")

# só a equipe
checar(cli.get("/api/admin/areas").status_code == 403, "cliente viu a calculadora")
checar(cli.post("/api/admin/areas/excel", json={"quadro": QUADRO}).status_code == 403, "cliente baixou o Excel")

# Excel: recusa quadro que não fecha; sai com as cores da casa
checar(adm.post("/api/admin/areas/excel", json={"quadro": ruim, "venda": 1}).status_code == 400, "exportou quadro que não fecha")
r = adm.post("/api/admin/areas/excel", json={"quadro": QUADRO, "venda": 35426700, "custo": "", "projeto": "Prancha teste"})
checar(r.status_code == 200 and r.data[:2] == b"PK", "o Excel não saiu")
if r.status_code == 200:
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
    textos = [str(c.value) for linha in ws.iter_rows() for c in linha if c.value is not None]
    checar(ws["A5"].fill.fgColor.rgb.endswith("002455"), "cabeçalho sem o azul da casa")
    checar(3542.67 in [c.value for linha in ws.iter_rows() for c in linha], "Excel sem a área construída")
    checar(10000 in [c.value for linha in ws.iter_rows() for c in linha], "Excel sem o R$/m² construído de 10.000")
    checar(any("privativa não consta" in t for t in textos), "Excel não avisa que a privativa não consta")
    checar(ws.page_setup.orientation == "portrait", "Excel não está em retrato")

if falhas:
    print("FALHAS:\n- " + "\n- ".join(falhas))
    sys.exit(1)
print("ok: calculadora de R$/m² (conferência, R$/m², Excel, acesso)")
